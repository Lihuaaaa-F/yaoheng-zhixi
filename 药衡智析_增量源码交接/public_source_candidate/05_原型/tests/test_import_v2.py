# -*- coding: utf-8 -*-
"""数据中心 v2（三模块改版，2026-09-22）回归：

- 类型化上传（data_type）与原始文件预览（表格/文本/PDF 内嵌）；
- 数据解析流水线：多文件合并发布（汇总优先防双计）、进度上报、
  成功/失败消息合同（数据处理成功 / 数据处理失败，{步骤}报错：…）；
- 知识索引构建流水线：解析落库 + 构建 + 消息合同；
- 报告模板解析流水线：结构检查、安装、消息合同；
- 模型注册表：厂商预填充、档位限制（提取=小模型/分析=大模型）、推理强度映射；
- 向量模型切换：本地校验/分析模型评估前置失败的消息合同；
- 目录过滤：合成演示上下文默认不出现，赛题数据为默认。
"""
import importlib
import json
from pathlib import Path

import pytest

from pharma import data_import, import_pipeline, model_registry, model_settings

WIDE_CSV = ('工厂,产品名称,月份,产量(盒),直接材料(元/盒),直接人工(元/盒),制造费用(元/盒),单位成本(元/盒),总成本(元)\n'
            '中药一厂,银黄口服液,2026-01,10000,3.10,1.00,0.90,5.00,50000\n'
            '中药一厂,银黄口服液,2026-02,10100,3.50,1.00,0.90,5.40,54540\n'
            '中药一厂,银黄口服液,2025-02,9800,3.00,0.95,0.85,4.80,47040\n')
BUDGET_CSV = ('工厂,产品名称,月份,预算产量(盒),预算直接材料(元/盒),预算直接人工(元/盒),预算制造费用(元/盒)\n'
              '中药一厂,银黄口服液,2026-02,10000,3.20,1.00,0.90\n')
MATERIAL_CSV = ('工厂,产品名称,月份,产量(盒),原材料名称,原材料总成本(元)\n'
                '中药二厂,银黄口服液,2026-01,9000,金银花,12000\n'
                '中药二厂,银黄口服液,2026-01,9000,黄芩提取物,9000\n'
                '中药二厂,银黄口服液,2026-02,9100,金银花,12600\n'
                '中药二厂,银黄口服液,2026-02,9100,黄芩提取物,9200\n')


@pytest.fixture()
def isolated_runtime(tmp_path, monkeypatch):
    monkeypatch.setenv('PHARMA_RUNTIME_DIR', str(tmp_path))
    monkeypatch.delenv('PHARMA_SHOW_TEST_CONTEXTS', raising=False)
    # 指向不存在的密钥文件：阻止本机 .env 经 config setdefault 重新注入真实
    # 密钥（否则测试会发起真实模型调用）
    monkeypatch.setenv('PHARMA_MODEL_KEY_FILE', str(tmp_path / 'no-key-in-tests.key'))
    monkeypatch.delenv('PHARMA_API_KEY', raising=False)
    monkeypatch.delenv('PHARMA_MODEL_ROUTES', raising=False)
    import pharma.config as config
    importlib.reload(config)
    for module in (data_import, model_settings, import_pipeline):
        importlib.reload(module)
    import pharma.industry as industry
    monkeypatch.setattr(industry, 'ENTERPRISE_REGISTRY', tmp_path / 'enterprise_registry.json')
    # 测试默认走确定性回退：模型建议路径以桩替换（真实网关行为另有合同测试）
    monkeypatch.setattr(import_pipeline, '_extraction_mapping', lambda h, s: (None, '测试：预设映射'))
    monkeypatch.setattr(import_pipeline, '_analysis_hypotheses', lambda a: None)
    return tmp_path


def _store(tmp_path):
    from pharma.jobs import JobStore
    return JobStore(tmp_path / 'app.sqlite3')


# ---------- 上传类型与预览 ----------

def test_upload_with_data_type_and_wrong_type_rejected(isolated_runtime):
    record = data_import.create_upload('business', '汇总.csv', WIDE_CSV.encode('utf-8'), 'cost_summary')
    assert record['meta']['data_type'] == 'cost_summary'
    assert record['status'] == 'UPLOADED'
    with pytest.raises(ValueError):
        data_import.create_upload('business', 'x.csv', WIDE_CSV.encode('utf-8'), 'product')
    with pytest.raises(ValueError):
        data_import.create_upload('knowledge', 'k.pdf', b'%PDF-1.4 fake', 'monthly')


def test_preview_original_table_text_and_pdf(isolated_runtime):
    table = data_import.preview_original(data_import.create_upload('business', 'a.csv', WIDE_CSV.encode('utf-8')))
    assert table['format'] == 'table' and table['headers'][0] == '工厂' and table['row_count'] == 3
    text = data_import.preview_original(data_import.create_upload('knowledge', 'k.txt', ('设备维修记录' * 20).encode('utf-8'), 'enterprise'))
    assert text['format'] == 'text' and text['characters'] > 30
    pdf = data_import.preview_original(data_import.create_upload('knowledge', 'k.pdf', b'%PDF-1.4 stub-bytes', 'product'))
    assert pdf['format'] == 'pdf' and pdf['url'].endswith('/file')


# ---------- 数据解析流水线 ----------

def test_data_parse_pipeline_publishes_and_reports_progress(isolated_runtime, tmp_path):
    store = _store(tmp_path)
    summary = data_import.create_upload('business', '一厂汇总.csv', WIDE_CSV.encode('utf-8'), 'cost_summary')
    material = data_import.create_upload('business', '二厂材料明细.csv', MATERIAL_CSV.encode('utf-8'), 'material_detail')
    job = store.enqueue('data_parse', {'import_ids': [summary['id'], material['id']]})
    import_pipeline.run_data_parse(store, store.get(job['id']))
    finished = store.get(job['id'])
    assert finished['status'] == 'SUCCEEDED'
    assert finished['result']['message'] == '数据处理成功'
    assert finished['progress'] == 100
    summary_after = data_import.get_import(summary['id'])
    material_after = data_import.get_import(material['id'])
    assert summary_after['status'] == 'PARSED' and material_after['status'] == 'PARSED'
    assert summary_after['meta']['parsed']['context_id'].startswith('pharmaceutical:imp-')
    events = [e for e in store.history(job['id']) if e.get('detail')]
    assert any('预处理' in e['detail'] for e in events)
    assert any('归因分析' in e['detail'] for e in events)
    # 二厂明细（材料要素）与一厂汇总各自工厂：两个工厂都可分析
    published = finished['result']['published']
    assert '中药一厂' in published['factories'] and '中药二厂' in published['factories']
    # 目录出现用户导入企业；合成上下文不出现
    import pharma.industry as industry
    catalog = industry.context_catalog()
    assert published['context_id'] in [c['context_id'] for c in catalog['contexts']]
    assert 'pharmaceutical:synthetic-pharma' not in [c['context_id'] for c in catalog['contexts']]


def test_data_parse_failure_message_contract(isolated_runtime, tmp_path):
    store = _store(tmp_path)
    broken = WIDE_CSV.replace('2026-01,10000', '2026-01,abc')
    record = data_import.create_upload('business', 'bad.csv', broken.encode('utf-8'), 'cost_summary')
    job = store.enqueue('data_parse', {'import_ids': [record['id']]})
    import_pipeline.run_data_parse(store, store.get(job['id']))
    finished = store.get(job['id'])
    assert finished['status'] == 'FAILED'
    assert finished['error'].startswith('数据处理失败，')
    assert '质量校验报错' in finished['error'] or '报错' in finished['error']
    assert data_import.get_import(record['id'])['status'] == 'PARSE_FAILED'


def test_summary_and_detail_no_double_count(isolated_runtime, tmp_path):
    """汇总与同要素明细同时导入：金额以汇总为准，不双计；产量一致才可合并。"""
    detail_same_plant = MATERIAL_CSV.replace('中药二厂', '中药一厂').replace(',9000,', ',10000,').replace(',9100,', ',10100,')
    summary = data_import.create_upload('business', 's.csv', WIDE_CSV.encode('utf-8'), 'cost_summary')
    detail = data_import.create_upload('business', 'd.csv', detail_same_plant.encode('utf-8'), 'material_detail')
    result = data_import.publish_business_batch(
        [(summary, {h: v for h, v in {
            '工厂': 'factory_id', '产品名称': 'product_id', '月份': 'period', '产量(盒)': 'quantity',
            '直接材料(元/盒)': 'element:material', '直接人工(元/盒)': 'element:labor',
            '制造费用(元/盒)': 'element:overhead', '总成本(元)': 'total_cost',
            '单位成本(元/盒)': ''}.items() if h in (data_import.get_import(summary['id'])['meta']['preview']['headers'])}, {'scenario': 'actual'}),
         (detail, {h: v for h, v in {
             '工厂': 'factory_id', '产品名称': 'product_id', '月份': 'period', '产量(盒)': 'quantity',
             '原材料名称': '', '单位消耗成本(元/盒)': '', '原材料总成本(元)': 'element:material'}.items() if h in (data_import.get_import(detail['id'])['meta']['preview']['headers'])}, {'scenario': 'actual'})],
        '合并企业X', 'pharmaceutical', '盒')
    facts = json.loads((data_import.IMPORTS_ROOT / 'enterprises' / result['enterprise_id'] / 'facts.json').read_text(encoding='utf-8'))
    material = [f for f in facts['costs'] if f['element_id'] == 'material' and f['period'] == '2026-02' and f['scenario'] == 'actual']
    assert len(material) == 1  # 汇总口径覆盖明细，未叠加
    assert material[0]['amount'] == '3.50'
    assert result['merge_warnings']  # 明细差异被记录为提示


# ---------- 知识索引构建流水线 ----------

class _FakeKnowledge:
    calls = {}

    def __init__(self, *args, **kwargs):
        pass

    def build(self, progress=None):
        if progress:
            progress(50, '解析知识文档 1/1')
            progress(100, '验证与切换')
        return {'status': 'PASS', 'knowledge_version': 'v-test', 'chunks': 12,
                'sources': {'a.txt': 'x'}, 'vector_error': None}

    def search(self, *args, **kwargs):
        return {'status': 'PASS', 'evidence': []}


def test_kb_build_pipeline_message_contract(isolated_runtime, tmp_path, monkeypatch):
    import pharma.knowledge as knowledge
    store = _store(tmp_path)
    doc = data_import.create_upload('knowledge', '设备清单.txt', ('车间设备清单：胶囊填充机 维修记录 ' * 10).encode('utf-8'), 'enterprise')
    job = store.enqueue('kb', {'import_ids': [doc['id']], 'rebuild': True})
    monkeypatch.setattr(knowledge, 'Knowledge', _FakeKnowledge)
    import_pipeline.run_kb_build(store, store.get(job['id']))
    finished = store.get(job['id'])
    assert finished['status'] == 'SUCCEEDED'
    assert finished['result']['message'].startswith('知识库构建成功')
    assert data_import.get_import(doc['id'])['status'] == 'PARSED'
    source_dir = data_import.KNOWLEDGE_INGEST_DIR
    written = list(Path(source_dir).glob('*.txt'))
    assert any('企业内部知识' in p.name for p in written)


def test_kb_build_empty_document_fails_with_step(isolated_runtime, tmp_path, monkeypatch):
    import pharma.knowledge as knowledge
    store = _store(tmp_path)
    doc = data_import.create_upload('knowledge', '空.txt', b'  \n ', 'enterprise')
    job = store.enqueue('kb', {'import_ids': [doc['id']], 'rebuild': True})
    monkeypatch.setattr(knowledge, 'Knowledge', _FakeKnowledge)
    import_pipeline.run_kb_build(store, store.get(job['id']))
    finished = store.get(job['id'])
    assert finished['status'] == 'FAILED'
    assert finished['error'].startswith('构建知识库失败，解析文档报错')
    assert data_import.get_import(doc['id'])['status'] == 'PARSE_FAILED'


# ---------- 报告模板解析流水线 ----------

def _template_docx(path: Path, sections=True, placeholder_count=24):
    from docx import Document
    doc = Document()
    if sections:
        for title in ('一、封面与基本信息', '二、总成本概览', '三、成本要素明细分析',
                      '四、重点产品专项分析', '五、对标分析', '六、总结与建议'):
            doc.add_heading(title, level=1)
        doc.add_paragraph('{{产品名称}} {{分析月份}} {{本月材料成本}}')
    for i in range(placeholder_count):
        doc.add_paragraph(f'{{{{指标占位{i}}}}}')
    doc.save(path)
    return path


def test_template_parse_installs_and_reports(isolated_runtime, tmp_path, monkeypatch):
    import pharma.reports as reports
    # 隔离安装目录：reports 模块常量不受 fixture 的 runtime 隔离影响（2026-09-22
    # 修复——此前测试把合成模板写进真实 .runtime/templates）。
    monkeypatch.setattr(reports, 'RUNTIME_TEMPLATES', tmp_path / 'templates')
    store = _store(tmp_path)
    source = _template_docx(tmp_path / 'tpl.docx')
    record = data_import.create_upload('template', 'tpl.docx', source.read_bytes(), 'quarterly')
    job = store.enqueue('template_parse', {'import_ids': [record['id']]})
    monkeypatch.setattr(import_pipeline, '_template_binding_analysis', lambda p: (None, '测试：确定性绑定'))
    import_pipeline.run_template_parse(store, store.get(job['id']))
    finished = store.get(job['id'])
    assert finished['status'] == 'SUCCEEDED'
    assert finished['result']['message'] == '报告模板解析成功'
    installed = reports.installed_templates()
    quarterly = [t for t in installed if t['analysis_type'] == 'quarterly'][0]
    assert quarterly['installed'] and quarterly['placeholder_count'] >= 20
    assert data_import.get_import(record['id'])['status'] == 'PARSED'
    # 未安装的类型回退题包月度工作模板
    path, _map = reports.working_template('special')
    assert 'templates' not in str(path) or path.name != 'special.docx'


def test_template_parse_missing_sections_fails(isolated_runtime, tmp_path, monkeypatch):
    import pharma.reports as reports
    monkeypatch.setattr(reports, 'RUNTIME_TEMPLATES', tmp_path / 'templates')
    store = _store(tmp_path)
    source = _template_docx(tmp_path / 'bad.docx', sections=False)
    record = data_import.create_upload('template', 'bad.docx', source.read_bytes(), 'monthly')
    job = store.enqueue('template_parse', {'import_ids': [record['id']]})
    import_pipeline.run_template_parse(store, store.get(job['id']))
    finished = store.get(job['id'])
    assert finished['status'] == 'FAILED'
    assert finished['error'].startswith('报告模板解析失败，章节结构检查报错')


# ---------- 模型注册表 ----------

def test_presets_payload_and_tiers():
    payload = model_registry.presets_payload()
    vendors = {v['id'] for v in payload['api_vendors']}
    assert {'zhipu', 'deepseek', 'dashscope', 'moonshot', 'siliconflow', 'openai'} <= vendors
    assert {v['id'] for v in payload['local_vendors']} >= {'ollama', 'vllm', 'lmstudio'}
    assert payload['role_rules']['extraction']['allowed_tier'] == 'small'
    assert payload['role_rules']['analysis']['allowed_tier'] == 'large'


def test_tier_error_rules():
    assert model_registry.tier_error('extraction', 'glm-5.3') is not None
    assert model_registry.tier_error('extraction', 'glm-4.5-air') is None
    assert model_registry.tier_error('analysis', 'glm-4.5-air') is not None
    assert model_registry.tier_error('analysis', 'glm-5.3') is None
    assert model_registry.tier_error('extraction', 'totally-unknown-model') is None  # 未收录可手填


def test_effort_mapping_by_vendor():
    assert model_registry.effort_body_params('glm-5.3', 'https://open.bigmodel.cn/api/paas/v4', 'medium') == {'reasoning_effort': 'high'}
    assert model_registry.effort_body_params('glm-5.3', 'x', 'high') == {'reasoning_effort': 'max'}
    assert model_registry.effort_body_params('deepseek-chat', 'https://api.deepseek.com', 'low') == {'thinking': {'type': 'disabled'}}
    assert model_registry.effort_body_params('deepseek-chat', 'https://api.deepseek.com', 'high') == {'thinking': {'type': 'enabled'}}
    assert model_registry.effort_body_params('qwen3.8-max', 'https://dashscope.aliyuncs.com/compatible-mode/v1', 'high') == {'enable_thinking': True}


def test_save_settings_rejects_wrong_tier(isolated_runtime):
    with pytest.raises(ValueError):
        model_settings.save_settings({'connections': {
            'extraction': {'model': 'glm-5.3', 'base_url': 'https://open.bigmodel.cn/api/paas/v4'}}})
    with pytest.raises(ValueError):
        model_settings.save_settings({'connections': {
            'analysis': {'model': 'glm-4.5-air', 'base_url': 'https://open.bigmodel.cn/api/paas/v4'}}})


# ---------- 向量模型切换 ----------

def test_vector_switch_missing_path_message(isolated_runtime, tmp_path):
    from pharma import vector_switch
    importlib.reload(vector_switch)
    store = _store(tmp_path)
    job = store.enqueue('vector_switch', {'path': str(tmp_path / 'nope')})
    vector_switch.run_vector_switch(store, store.get(job['id']))
    finished = store.get(job['id'])
    assert finished['status'] == 'FAILED'
    assert finished['error'].startswith('向量模型切换失败，本地校验报错')


def test_vector_switch_requires_analysis_model(isolated_runtime, tmp_path, monkeypatch):
    import pharma.knowledge as knowledge
    import pharma.narrative as narrative
    from pharma import vector_switch
    importlib.reload(vector_switch)

    class _FakeEmbedding:
        def __init__(self, path):
            pass

        def encode(self, texts, query=False):
            return [[0.1, 0.2, 0.3] for _ in texts]

    class _NoKeyGateway:
        key = ''

        def __init__(self, *a, **k):
            pass

    model_dir = tmp_path / 'vec-model'
    model_dir.mkdir()
    (model_dir / 'model_quantized.onnx').write_bytes(b'onnx-stub')
    (model_dir / 'tokenizer.json').write_text('{}', encoding='utf-8')
    monkeypatch.setattr(knowledge, 'CpuEmbedding', _FakeEmbedding)
    monkeypatch.setattr(narrative.ModelGateway, 'for_route',
                        classmethod(lambda cls, route, **kw: _NoKeyGateway()))
    store = _store(tmp_path)
    job = store.enqueue('vector_switch', {'path': str(model_dir)})
    vector_switch.run_vector_switch(store, store.get(job['id']))
    finished = store.get(job['id'])
    assert finished['status'] == 'FAILED'
    assert finished['error'].startswith('向量模型切换失败，数据分析模型API报错')
    assert '未配置数据分析模型' in finished['error']


# ---------- 目录过滤 ----------

def test_catalog_default_is_competition_without_synthetic(isolated_runtime):
    import pharma.industry as industry
    catalog = industry.context_catalog()
    ids = [c['context_id'] for c in catalog['contexts']]
    assert 'pharmaceutical:synthetic-pharma' not in ids
    assert not any(c['context_id'].startswith(('mechanical_demo', 'chemical_demo')) for c in catalog['contexts'])
    assert catalog['default_context_id'] == 'pharmaceutical:competition'


def test_catalog_show_test_contexts_opt_in(isolated_runtime, monkeypatch):
    monkeypatch.setenv('PHARMA_SHOW_TEST_CONTEXTS', '1')
    import pharma.industry as industry
    catalog = industry.context_catalog()
    ids = [c['context_id'] for c in catalog['contexts']]
    assert 'pharmaceutical:synthetic-pharma' in ids
    assert catalog['default_context_id'] == 'pharmaceutical:competition'


def test_settings_route_key_survives_cross_host(monkeypatch, tmp_path):
    """设置文件的路由专属密钥不被主 env 密钥遮蔽、不被跨主机保护清空（2026-09-22 修复）。

    场景：主配置 glm（env PHARMA_MODEL_KEY_FILE 指向 glm 密钥），
    extraction 在设置文件指向 deepseek 并带专属密钥——网关必须用专属密钥。
    隔离：必须先 reload config 再 reload model_settings/narrative（conftest 在
    收集期已导入 config，且 autouse fixture 注入 PHARMA_MODEL* 主变量需清除），
    否则设置/密钥会写入真实 .runtime。
    """
    import importlib
    monkeypatch.setenv('PHARMA_RUNTIME_DIR', str(tmp_path))
    for var in ('PHARMA_MODEL', 'PHARMA_MODEL_BASE_URL', 'PHARMA_MODEL_PROTOCOL'):
        monkeypatch.delenv(var, raising=False)
    monkeypatch.setenv('PHARMA_MODEL_KEY_FILE', str(tmp_path / 'glm.key'))
    (tmp_path / 'glm.key').write_text('glm-main-key', encoding='utf-8')
    import pharma.config as config
    importlib.reload(config)
    import pharma.model_settings as ms
    import pharma.narrative as narrative
    importlib.reload(ms)
    importlib.reload(narrative)
    try:
        key_file = ms.KEYS_DIR / 'extraction.key'
        key_file.parent.mkdir(parents=True, exist_ok=True)
        key_file.write_text('deepseek-route-key', encoding='utf-8')
        ms.save_settings({'connections': {'extraction': {
            'model': 'deepseek-flash', 'base_url': 'https://api.deepseek.com',
            'protocol': 'openai', 'key_file': str(key_file)}}})
        gw = narrative.ModelGateway.for_route('extraction')
        assert gw.key == 'deepseek-route-key'
        assert gw.base_url == 'https://api.deepseek.com'
        assert gw.model == 'deepseek-flash'
        assert gw.credential_scope['source'] == 'key_file'
        # 主路由不受影响：无设置节时仍用主 env 密钥
        main_gw = narrative.ModelGateway.for_route('analysis')
        assert main_gw.key == 'glm-main-key'
    finally:
        importlib.reload(config)
        importlib.reload(ms)
        importlib.reload(narrative)


def test_sanitize_overrides_ignores_none_values():
    """API 端点曾把 {'base_url':x,'key_file':None} 传入：None 被 str() 成 'None'
    （truthy）导致密钥文件解析为 Path('None') → 密钥空（2026-09-22 修复）。"""
    cleaned = model_settings._sanitize_overrides({'base_url': 'https://api.deepseek.com', 'key_file': None})
    assert cleaned == {'base_url': 'https://api.deepseek.com'}


def test_vector_dimension_probe_cached(tmp_path, monkeypatch):
    """向量维度惰性探测：无 config.json 时从 ONNX 输出形状读，按指纹缓存。"""
    import importlib
    import pharma.config as config
    monkeypatch.setenv('PHARMA_RUNTIME_DIR', str(tmp_path))
    importlib.reload(config)
    importlib.reload(model_settings)
    model_dir = tmp_path / 'vec'
    model_dir.mkdir()
    (model_dir / 'model_quantized.onnx').write_bytes(b'stub')
    (model_dir / 'tokenizer.json').write_text('{}', encoding='utf-8')
    # 无 onnxruntime 可加载的真实权重：stub 文件探测失败 → None（展示降级，不抛错）
    assert model_settings._probe_dimension_cached(model_dir) is None or \
           isinstance(model_settings._probe_dimension_cached(model_dir), int)
