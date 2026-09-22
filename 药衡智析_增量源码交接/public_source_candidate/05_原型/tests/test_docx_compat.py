"""Word 兼容出厂校验（docx_compat）与模板命名空间前缀保留的回归测试。

2026-09-22 事故：模板手术用 xml.etree 重序列化整份 document.xml，前缀被
改名 ns0:/ns1:，根元素 mc:Ignorable 引用未声明前缀，MS Word 报"文件可能
已经损坏"（wdmain11.chm 25272）拒开；LibreOffice 宽容放行所以 PDF 与机器
审计全绿。本组测试钉住四件事：①校验器必须识别该损坏形态；②干净文件不
误报；③模板安装/规范化管线产物保留 w:/w14 前缀并通过出厂校验；④渲染
产出（含参照报告分支）通过出厂校验。
"""
import zipfile
import pytest
from pathlib import Path
from docx import Document
from pharma.docx_compat import validate_word_compat


def _placeholder_docx(path) -> Path:
    """python-docx 默认模板自带 mc:Ignorable="w14 wp14"，正是事故触发面。"""
    doc = Document()
    doc.add_heading('一、封面与基本信息', level=1)
    doc.add_paragraph('{{产品名称}} {{分析月份}}')
    doc.save(path)
    return Path(path)


def test_validator_rejects_dangling_ignorable_prefix(tmp_path):
    """复刻旧事故路径（标准库 ET 往返）→ 校验器必须拦截并点名悬空前缀。"""
    from xml.etree import ElementTree as ET
    source = _placeholder_docx(tmp_path / 'src.docx')
    out = tmp_path / 'corrupt.docx'
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            raw = zin.read(info.filename)
            if info.filename == 'word/document.xml':
                raw = ET.tostring(ET.fromstring(raw), encoding='utf-8', xml_declaration=True)
            zout.writestr(info, raw)
    xml = out.read_bytes() and zipfile.ZipFile(out).read('word/document.xml').decode('utf-8')
    assert 'xmlns:ns0' in xml[:500]  # 事故形态确实形成（前缀被改名）
    with pytest.raises(ValueError, match='WORD_COMPAT_CHECK_FAILED') as excinfo:
        validate_word_compat(out)
    assert 'w14' in str(excinfo.value)


def test_validator_passes_clean_documents(tmp_path):
    path = _placeholder_docx(tmp_path / 'clean.docx')
    assert validate_word_compat(path)['status'] == 'PASS'


def test_install_template_preserves_prefixes_and_passes_gate(tmp_path, monkeypatch):
    import pharma.reports as reports
    monkeypatch.setattr(reports, 'RUNTIME_TEMPLATES', tmp_path / 'templates')
    source = _placeholder_docx(tmp_path / 'tpl.docx')
    result = reports.install_template(source, 'monthly')
    installed = tmp_path / 'templates' / 'monthly.docx'
    assert installed.is_file()
    assert result['placeholder_count'] >= 1
    assert validate_word_compat(installed)['status'] == 'PASS'
    xml = zipfile.ZipFile(installed).read('word/document.xml').decode('utf-8')
    assert 'mc:Ignorable="w14 wp14"' in xml and '<w:p' in xml
    assert 'xmlns:ns0' not in xml[:600]  # lxml 保留原前缀，不再产生 ns0 改名


def test_normalize_template_output_is_word_compatible(tmp_path):
    from pharma.reports import normalize_template
    out = tmp_path / 'work.docx'
    normalize_template(out, tmp_path / 'map.json')
    assert validate_word_compat(out)['status'] == 'PASS'
    xml = zipfile.ZipFile(out).read('word/document.xml').decode('utf-8')
    assert '<w:tbl' in xml and 'w14:paraId' in xml
    assert 'xmlns:ns0' not in xml[:600]


def test_reference_render_output_passes_word_compat(tmp_path):
    from pharma.reports import render_docx
    from pharma.industry import analyze_reference
    snapshot = analyze_reference('pharmaceutical:synthetic-pharma', month='2026-06')
    snapshot['trend'] = []
    narrative = {'findings': [{'section': 'materials', 'claim_type': 'insufficient_evidence',
                               'origin': 'model', 'rendered_text': '合成缺证说明：未提供领料记录，不能归因。'}]}
    path = tmp_path / 'generic.docx'
    render_docx(snapshot, narrative, {'evidence': []}, path)
    assert validate_word_compat(path)['status'] == 'PASS'
