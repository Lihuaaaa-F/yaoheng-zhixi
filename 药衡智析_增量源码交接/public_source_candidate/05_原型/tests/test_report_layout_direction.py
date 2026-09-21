"""Synthetic regressions for numeric grouping and report comparison labels."""
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest
from pharma.industry import analyze_reference
from pharma.reports import build_bindings, style_reader, layout_text


def test_numeric_unit_group_preserves_runs_and_bookmark():
    doc=Document();p=doc.add_paragraph()
    mark=OxmlElement('w:bookmarkStart');mark.set(qn('w:id'),'11');mark.set(qn('w:name'),'metric_position');p._p.append(mark)
    p.add_run('本期变动 0.2');p.add_run('2 元/盒；产量 1200.00 件。')
    original=p.text
    style_reader(doc)
    assert '0\ufeff.\ufeff2\ufeff2' in p.text
    assert layout_text(p.text) == original
    assert p._p.xpath('.//w:bookmarkStart[@w:name="metric_position"]')
    once=p.text;style_reader(doc);assert p.text==once


def test_style_reader_preserves_template_content_and_styles_additions():
    """四轮E项：style_reader 不再覆盖模板内容样式；只装饰图题与目录行。"""
    doc=Document();heading=doc.add_heading('总成本概览',1)
    source=doc.add_paragraph('证据来源');doc.add_paragraph('合成证据条目')
    cap=doc.add_paragraph('图｜测试图题');toc=doc.add_paragraph(chr(0x3000)+'一、测试')
    heading.runs[0].font.size=Pt(15)
    style_reader(doc)
    assert heading.runs[0].font.size==Pt(15)  # 模板/普通内容样式不被覆盖
    assert not (source.paragraph_format.first_line_indent or 0)
    assert cap.alignment is not None and cap.runs[0].font.size==Pt(8.5)
    assert toc.runs[0].font.size==Pt(10.5)


@pytest.mark.parametrize('left,right', [('示范工厂A','示范工厂B'),('示范工厂B','示范工厂A')])
def test_comparison_binding_follows_actual_direction(left,right):
    s=analyze_reference('pharmaceutical:synthetic-pharma',month='2026-06')
    comparison={'left':left,'right':right,'direction':left+'−'+right+'，以'+right+'为分母',
                'elements':[{'name':'直接材料','delta':'-2','contribution':'100','unit':'元/盒'}]}
    text=build_bindings(s,{},comparison)['差异结构拆解分析']
    assert comparison['direction'] in text
    assert '二厂减一厂' not in text


@pytest.mark.parametrize('left,right', [('示范工厂A','示范工厂B'),('示范工厂B','示范工厂A')])
def test_chart_legend_and_caption_follow_comparison(tmp_path, monkeypatch, left, right):
    from pharma.reports import add_reader_charts
    from matplotlib.axes import Axes
    s=analyze_reference('pharmaceutical:synthetic-pharma',month='2026-06')
    c={'left':left,'right':right,'direction':left+'−'+right+'，以'+right+'为分母',
       'elements':[{'name':'直接材料','left':'4','right':'6','delta':'-2','unit':'元/盒'}]}
    doc=Document();doc.add_paragraph('2.2 合成成本结构')
    anchors={key:doc.add_paragraph(key) for key in ('近6个月成本趋势表格','对标差异表格')}
    seen=[];original=Axes.legend
    def legend(self,*args,**kwargs):
        result=original(self,*args,**kwargs)
        seen.extend(t.get_text() for t in result.get_texts());return result
    monkeypatch.setattr(Axes,'legend',legend)
    add_reader_charts(doc,s,c,anchors,tmp_path/'report.docx')
    assert seen == [right,left]
    assert any(c['direction'] in p.text for p in doc.paragraphs)
    assert all('二厂−一厂' not in p.text for p in doc.paragraphs)


def test_template_identity_sanitized_by_role_without_private_literals():
    from pharma.reports import sanitize_template_identity
    doc=Document();table=doc.add_table(rows=2,cols=2)
    table.cell(0,0).text='编制人';table.cell(0,1).text='合成人名甲'
    table.cell(1,0).text='审核人';table.cell(1,1).text='合成人名乙'
    p=doc.add_paragraph();box=OxmlElement('w:txbxContent');inner=OxmlElement('w:p')
    run=OxmlElement('w:r');text=OxmlElement('w:t');text.text='独立合成公司水印'
    run.append(text);inner.append(run);box.append(inner);p._p.append(box)
    sanitize_template_identity(doc)
    assert table.cell(0,1).text=='财务部成本会计'  # 模板md固定值（2026-09-21三轮：模板为准）
    assert table.cell(1,1).text=='财务总监'  # 模板md固定值
    assert text.text=='药衡智析 · 成本分析'
    assert '合成人名' not in doc.element.xml


def test_footer_replaces_legacy_textbox_and_preserves_explicit_page_roles():
    from pharma.reports import rebuild_report_footer
    doc=Document();footer=doc.sections[0].footer
    p=footer.paragraphs[0];p.add_run('错误品牌16')
    box=OxmlElement('w:txbxContent');p._p.append(box)
    rebuild_report_footer(doc)
    assert not footer._element.xpath('.//*[local-name()="txbxContent"]')
    assert [x.get(qn('w:instr')) for x in footer._element.xpath('.//w:fldSimple')]==['PAGE','NUMPAGES']
    all_text=''.join(x.text or '' for x in footer._element.iter(qn('w:t')))
    assert all_text=='第1页共1页'  # 模板页脚格式（四轮E项）
    rebuild_report_footer(doc)
    assert len(footer.paragraphs)==1


def test_number_unit_space_is_nonbreaking_and_short_source_block_stays_together():
    doc=Document();p=doc.add_paragraph('每盒变动 0.18 元。')
    source=doc.add_paragraph('来源与审核说明');entry=doc.add_paragraph('独立合成文档，第1页')
    final=doc.add_paragraph('人工归因评分待真实评审。')
    style_reader(doc)
    assert '0.18\u00a0元' in p.text.replace('\ufeff','')
    assert source.paragraph_format.keep_with_next and entry.paragraph_format.keep_with_next
    assert not final.paragraph_format.keep_with_next


def test_model_missing_evidence_text_is_bound_to_each_actual_section(tmp_path):
    from pharma.reports import insert_element_analysis, explanation_presence
    snapshot=analyze_reference('pharmaceutical:synthetic-pharma',month='2026-06')
    findings=[{'section':section,'claim_type':'insufficient_evidence','origin':'model',
               'rendered_text':f'独立合成{section}缺证说明：缺少签署原始记录，不能确定原因。'}
              for section in ('materials','labor','overhead','benchmark')]
    findings.append({'section':'labor','claim_type':'recommendation','origin':'model','rendered_text':'不得混入正文的唯一行动指令'})
    narrative={'findings':findings}
    bindings=build_bindings(snapshot,narrative)
    doc=Document();doc.add_heading('3.1 直接材料分析',2);doc.add_paragraph(bindings['材料成本归因分析文本'])
    doc.add_heading('3.2 直接人工分析',2);doc.add_heading('3.3 制造费用分析',2)
    doc.add_heading('四、重点产品专项分析',1);doc.add_heading('5.3 差异归因分析',2)
    doc.add_paragraph(bindings['差异归因分析文本']);doc.add_heading('六、总结与建议',1)
    insert_element_analysis(doc,snapshot,bindings)
    path=tmp_path/'sections.docx';doc.save(path)
    result=explanation_presence(path,narrative)
    assert result['explanation_bindings_checked']==4 and not result['explanation_binding_failures']
    all_text='\n'.join(p.text for p in doc.paragraphs)
    for finding in findings[:4]:assert all_text.count(finding['rendered_text'])==1
    assert findings[4]['rendered_text'] not in all_text
    # Moving the only valid text to the recommendation table cannot pass.
    target=next(p for p in doc.paragraphs if findings[1]['rendered_text'] in p.text)
    target.text=''
    doc.add_table(rows=1,cols=1).cell(0,0).text=findings[1]['rendered_text']
    doc.save(path)
    assert explanation_presence(path,narrative)['explanation_binding_failures']==[
        {'finding':1,'section':'labor','reason':'MODEL_EXPLANATION_MISSING_FROM_BODY'}]


def test_verify_docx_fails_when_model_explanation_is_removed(tmp_path):
    from pharma.reports import render_docx, verify_docx
    snapshot=analyze_reference('pharmaceutical:synthetic-pharma',month='2026-06');snapshot['trend']=[]
    finding={'section':'materials','claim_type':'insufficient_evidence','origin':'model',
             'rendered_text':'独立合成材料缺证：未提供实际领料记录，不能归因。'}
    narrative={'findings':[finding]};path=tmp_path/'generic.docx'
    assert render_docx(snapshot,narrative,{'evidence':[]},path)['verification']['explanation_bindings_checked']==1
    doc=Document(path)
    next(p for p in doc.paragraphs if p.text==finding['rendered_text']).text='正文缺失'
    doc.save(path)
    result=verify_docx(path,snapshot,narrative=narrative)
    assert result['status']=='FAIL' and len(result['explanation_binding_failures'])==1


def test_template_prose_rewrite_never_changes_inserted_model_text_or_split_slots(tmp_path):
    from pharma import reports
    doc=Document();p=doc.add_paragraph()
    p.add_run('与模拟乙厂；本月 {{本');p.add_run('月产量}}：{{解释}}')
    original='需核查模拟甲厂与模拟乙厂的本月成本记录。'
    reports.rewrite_template_prose(p,[('与模拟乙厂','模拟甲厂−模拟乙厂'),('本月','本季度')])
    reports.replace_text_nodes(list(p._p.iter('{'+reports.W+'}t')),{'本月产量':'120','解释':original})
    assert p.text=='模拟甲厂−模拟乙厂；本季度 120：'+original
    assert '{{' not in p.text
    assert len(p.runs)==2


def test_footer_body_clearance_and_full_period_label():
    from pharma.reports import report_period_label, benchmark_precision
    doc=Document();doc.add_paragraph('独立合成正文');style_reader(doc)
    section=doc.sections[0]
    # 四轮E项：模板节边距不再被覆盖（默认 2.54cm 保持原样）
    assert section.bottom_margin.cm == pytest.approx(2.54,abs=.01)
    assert section.bottom_margin.pt-section.footer_distance.pt > 30
    context={'month':'2026-06','analysis_type':'quarterly','period':{'start':'2026-04','end':'2026-06'}}
    assert report_period_label(context)=='2026-04 至 2026-06'
    assert benchmark_precision(context)==4
    assert report_period_label({'period':{'start':'2026-06','end':'2026-06'}})=='2026-06'


def test_quarterly_benchmark_preserves_small_difference(tmp_path):
    from pharma.reference_report import render,verify
    from pharma.reports import layout_text
    snapshot=analyze_reference('mechanical_demo:synthetic-mechanical',month='2026-06',analysis_type='quarterly')
    snapshot['trend']=[]
    comparison={'left':'合成工厂甲','right':'合成工厂乙','elements':[
        {'key':'material','name':'合成材料','left':'8.1234','right':'8.1212','delta':'0.0022'}]}
    path=tmp_path/'quarter.docx'
    result=render(snapshot,{'findings':[]},{'evidence':[]},path,comparison)
    assert result['status']=='PASS'
    assert verify(path,snapshot,comparison)['status']=='PASS'
    rows=[[layout_text(c.text) for c in r.cells] for t in Document(path).tables for r in t.rows]
    assert ['合成材料','8.1234','8.1212','0.0022'] in rows


def test_actual_pdf_number_unit_groups_and_footer_clearance(tmp_path):
    import shutil,re
    import pymupdf
    from pharma.reports import convert_pdf
    if not shutil.which('libreoffice'):pytest.skip('LibreOffice unavailable')
    doc=Document()
    # Independent synthetic fixture moves the same amount across line ends.
    for n in range(31,43):
        doc.add_paragraph('合'*n+' 12345.67 元；合成产量 345.00 件。')
    style_reader(doc);path=tmp_path/'group.docx';doc.save(path)
    converted=convert_pdf(path)
    assert converted['status']=='PASS',converted
    pdf=pymupdf.open(path.with_suffix('.pdf'))
    occurrences=0
    for page in pdf:
        lines=[line for block in page.get_text('dict')['blocks'] for line in block.get('lines',[])]
        body=[];foot=[]
        for line in lines:
            text=''.join(s['text'] for s in line['spans'])
            if re.match(r'^第\s*\d+\s*页',text):foot.append(line['bbox'])
            else:body.append(line['bbox'])
            if '12345.67' in text:
                occurrences+=1
                assert '12345.67 元' in text.replace('\u00a0',' ')
        assert min(b[1] for b in foot)-max(b[3] for b in body)>10
    assert occurrences==12
