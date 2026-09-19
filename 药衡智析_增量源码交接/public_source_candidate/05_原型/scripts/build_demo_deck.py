"""Build the existing ten-slide deck from one explicitly bound acceptance run.

Native PPTX rendering is separate. An optional reading preview is written only to
_reading_preview; it never replaces a LibreOffice rendering of the actual slides.
"""
from pathlib import Path
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import argparse
import json

APP = Path(__file__).resolve().parents[1]
ROOT = APP.parent
SCENARIOS = ('S1', 'S2', 'S3')


def read_json(path):
    return json.loads(path.read_text(encoding='utf-8'))


def resolve_path(value, manifest_path, reports_dir=None):
    path = Path(value)
    candidates = [path] if path.is_absolute() else [manifest_path.parent/path, ROOT/path, APP/path]
    if reports_dir and not path.is_absolute():
        candidates.insert(0, reports_dir/path)
    return next((p.resolve() for p in candidates if p.is_file()), None)


def amount(value):
    if value is None:
        return '缺失'
    try:
        number = Decimal(str(value))
        return format(number.quantize(Decimal('0.0001')), 'f').rstrip('0').rstrip('.') if number.is_finite() else '缺失'
    except InvalidOperation:
        return '缺失'


def load_run(manifest_path=None, reports_dir=None):
    if not manifest_path:
        return {'run_id': '未绑定运行草稿', 'scenarios': [], 'human_review': 'PENDING'}, {}
    manifest_path = manifest_path.resolve()
    manifest = read_json(manifest_path)
    if not manifest.get('run_id') or not isinstance(manifest.get('scenarios'), list):
        raise ValueError('MANIFEST_RUN_AND_SCENARIOS_REQUIRED')
    if len({row['id'] for row in manifest['scenarios']}) != len(manifest['scenarios']):
        raise ValueError('DUPLICATE_SCENARIO')
    loaded = {}
    for row in manifest['scenarios']:
        sid = row['id']; item = {'row': row, 'snapshot': {}, 'artifacts': {}, 'narrative': {}}
        for fmt in ('docx', 'pdf'):
            binding = (row.get('artifacts') or {}).get(fmt) or {}
            candidate = resolve_path(binding['path'], manifest_path, reports_dir) if binding.get('path') else None
            if candidate is None and reports_dir:
                candidate = next((p for p in (reports_dir/sid/f'report.{fmt}', reports_dir/f'{sid}.{fmt}') if p.is_file()), None)
            if candidate:
                # Directory names do not establish provenance. Every report must
                # match the artifact hash captured for this scenario and job.
                if not binding.get('sha256'):
                    raise ValueError(f'UNBOUND_REPORT:{sid}:{fmt}')
                if sha256(candidate.read_bytes()).hexdigest() != binding['sha256']:
                    raise ValueError(f'REPORT_HASH_MISMATCH:{sid}:{fmt}')
                item['artifacts'][fmt] = candidate
        paths = []
        for key in ('snapshot_path', 'result_path'):
            if row.get(key):
                candidate = resolve_path(row[key], manifest_path, reports_dir)
                if candidate: paths.append(candidate)
        snapshot_binding = (row.get('artifacts') or {}).get('snapshot') or {}
        if snapshot_binding.get('path'):
            candidate = resolve_path(snapshot_binding['path'], manifest_path, reports_dir)
            if candidate: paths.append(candidate)
        if reports_dir:
            paths += [reports_dir/sid/'job_result.json', reports_dir/sid/'snapshot.json', reports_dir/f'{sid}_snapshot.json']
        for candidate in paths:
            if not candidate.is_file(): continue
            data = read_json(candidate)
            body = data.get('result', data)
            snapshot = body.get('snapshot', body)
            if snapshot.get('snapshot_id') != row.get('snapshot_id') or not row.get('snapshot_id'):
                raise ValueError(f'SNAPSHOT_BINDING_MISMATCH:{sid}')
            item['snapshot'] = snapshot
            item['narrative'] = body.get('narrative') or {}
            item['snapshot_path'] = candidate
            break
        loaded[sid] = item
        screenshot = (row.get('artifacts') or {}).get('screenshot') or {}
        if screenshot.get('path'):
            candidate = resolve_path(screenshot['path'], manifest_path, reports_dir)
            if candidate:
                if not screenshot.get('sha256') or sha256(candidate.read_bytes()).hexdigest() != screenshot['sha256']:
                    raise ValueError(f'SCREENSHOT_HASH_MISMATCH:{sid}')
                item['screenshot'] = candidate
    receipts = {}
    for name in ('model_live', 'retrieval', 'rpa', 'browser'):
        value = (manifest.get('receipts') or {}).get(name)
        if not isinstance(value, str): continue
        path = resolve_path(value, manifest_path)
        if not path: continue
        receipt = read_json(path)
        if receipt.get('run_id') != manifest['run_id']:
            raise ValueError(f'RECEIPT_RUN_MISMATCH:{name}')
        if manifest.get('attempt_id') and receipt.get('attempt_id') != manifest['attempt_id']:
            raise ValueError(f'RECEIPT_ATTEMPT_MISMATCH:{name}')
        receipts[name] = receipt
    for sid, item in loaded.items():
        for name, receipt in receipts.items():
            row = next((r for r in receipt.get('scenarios', []) if r.get('id') == sid), None)
            if row:
                for key in ('job_id', 'snapshot_id', 'attempt_id'):
                    if item['row'].get(key) and row.get(key) != item['row'][key]:
                        raise ValueError(f'RECEIPT_SCENARIO_MISMATCH:{name}:{sid}:{key}')
                item[name] = row
    return manifest, loaded


def metric_line(snapshot, key, label):
    metric = (snapshot.get('metrics') or {}).get(key) or {}
    return f'{label}：{amount(metric.get("value"))} {metric.get("unit") or ""}'.strip()


def scenario_slide(sid, item):
    snapshot = item.get('snapshot') or {}
    narrative = item.get('model_live') or item.get('narrative') or {}
    identity = narrative.get('identity') or narrative.get('model_identity') or {}
    verified = bool(narrative.get('model_live')) and identity.get('status') == 'VERIFIED'
    model = ' / '.join(identity.get('returned') or []) or identity.get('returned_model') or identity.get('model') or narrative.get('model')
    if verified and model:
        mode = f'模型解释：{model}，身份已核验'
    elif narrative.get('generation_mode'):
        mode = '解释来源：确定性事实与待核验归因；本轮模型身份待核验'
    else:
        mode = '模型解释：未完成本轮实调核验'
    docs = item.get('artifacts') or {}
    title = f'{sid} · {snapshot.get("product") or "场景结果待生成"}'
    lines = [f'{snapshot.get("factory") or "工厂待绑定"} · {snapshot.get("month") or "期间待绑定"}',
             metric_line(snapshot, 'total_cost', '总成本'), metric_line(snapshot, 'unit_cost', '单位成本'),
             metric_line(snapshot, 'mom', '环比'),
             f'报告：Word {"已核验" if "docx" in docs else "待完成"} / PDF {"已核验" if "pdf" in docs else "待完成"}',
             mode, '专业原因、可读性及 0—5 分评分：待真人署名']
    note = {'scenario': sid, 'binding': item.get('row'), 'snapshot': str(item.get('snapshot_path', '')),
            'reports': {k: str(v) for k, v in docs.items()}, 'model_evidence': narrative}
    return {'title': title, 'lines': lines, 'pdf': docs.get('pdf'), 'screenshot': item.get('screenshot'), 'notes': note}


def build_slides(manifest, loaded):
    scenario_items = [loaded.get(sid, {}) for sid in SCENARIOS]
    report_count = sum(len(item.get('artifacts') or {}) == 2 for item in scenario_items)
    rpa = [item.get('rpa') or {} for item in scenario_items]
    sent = sum(row.get('status') == 'PASS' and row.get('notification') == 'SIMULATED_SENT' for row in rpa)
    def slide(title, lines, **kw): return {'title': title, 'lines': lines, **kw}
    overview = []
    for sid, item in zip(SCENARIOS, scenario_items):
        snapshot = item.get('snapshot') or {}
        overview.append(f'{sid} · {snapshot.get("product") or "场景待绑定"} · {snapshot.get("month") or "期间待绑定"}')
        overview.append('  ' + metric_line(snapshot, 'unit_cost', '单位成本'))
    return [
        slide('药衡智析 · 从成本异常到可核查行动', ['制药成本分析、证据引用与 Word / PDF 报告',
            f'本轮三个核心场景：{report_count}/3 组报告文件已核验',
            f'本地模拟 RPA 送达：{sent}/3；不等于整改完成',
            '真人专业归因、可读性与版式评审仍待完成'], notes={'run_id': manifest['run_id'], 'revision': manifest.get('revision') or manifest.get('commit')}),
        slide('三场景结果 · 先看数值，再核对证据', overview, notes={'source': '本轮绑定分析快照；缺失数据不补零'}),
        *[scenario_slide(sid, item) for sid, item in zip(SCENARIOS, scenario_items)],
        slide('指标证据能重算，缺证结论有边界', ['金额使用 Decimal；产量独立去重，不随成本明细累加',
            '季度单位成本 = Σ成本 / Σ可比产量，不能直接平均月单价',
            '机械合成反例：材料 15 对 30 元/件 → (15−30)/30 = −50%',
            '差额、差异率、贡献度分别保存分子、分母与双方来源',
            '环比、同比、预算和跨厂分别核对期间；缺失和零分母保留无定义'],
            notes={'tests': 'tests/test_industry.py', 'scope': '独立合成反例，不是比赛原始数据或行业基准'}),
        slide('任务闭环 · 确认载荷与送达回执逐项核对', [
            '报告发现 → 编辑负责人 / 行动 → 用户确认 → 幂等 outbox',
            f'本轮核心场景本地模拟送达：{sent}/3；其他场景见技术附件',
            '自动演示只访问回环地址，并在确认前核验服务端模拟模式',
            '同一任务编号但行动或收件人不同，不承认送达',
            '查询超时保留此前已证实的送达证据；真人责任确认待本人填写'],
            notes={'rpa_receipts': rpa, 'notice': '演示自动确认不冒充真人确认或真实微信发送'}),
        slide('加分原型 · 用小规模对照说明实际能力', [
            '图谱：只扩展 BM25 词项；向量与重排仍使用原问题，可关闭',
            '固定三个合成问题：关闭召回 2/3，开启 3/3，产品范围不串用',
            '该开发样本不证明普遍增益；多模型实调以本轮身份回执为准',
            '预测反例：10 / 20 / 30 → 下期 40；末次值基线为 30',
            '常量不漂移；缺月拒绝连续外推；波动范围未经覆盖率验证'],
            notes={'tests': ['tests/test_purpose_retrieval.py::test_fixed_graph_toggle_comparison_uses_same_scope_and_no_model', 'tests/test_forecasting.py'], 'forecast': '样本内一步误差非独立留出检验；不宣称专业预测已验收'}),
        slide('行业迁移 · 共用核心，保留专属边界', [
            '机械参考包：3600 元 / 120 件 = 30 元/件；单位机时 0.5 小时/件',
            '化工参考包：2400 元 / 200 kg = 12 元/kg；单位能耗 3 kWh/kg',
            '制药、机械、化工双产品合成回归核对数值、单位与来源隔离',
            '上述数字是独立合成假设，不是行业基准或全制造业适配证明',
            '下一步由行业队友核对实际数据、政策、知识及联副产品 / 在制品口径'],
            notes={'tests': 'tests/test_industry.py', 'architecture': '现有模块化单体与受信行业包，无任意代码插件'}),
        slide('交付与待评 · 自动化证据和真人判断分开', [
            f'三场景 Word / PDF：{report_count}/3；本地模拟送达：{sent}/3',
            '打开交付目录中的 S1、S2、S3 报告，与数字和引用来源逐项核对',
            '真人分别填写专业归因、可读性和版式意见，再署名评分',
            '模型、检索、数值、浏览器与 RPA 的结果分别保存在技术附件',
            '缺少真人评分时保留待评状态，不宣称比赛整体验收或获奖水平'],
            notes={'run_id': manifest['run_id'], 'human_review': manifest.get('human_review', 'PENDING'), 'competition_ready': False})]


def write_deck(slides, output_dir, reading_preview=False):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    import fitz
    output_dir.mkdir(parents=True, exist_ok=True)
    pres = Presentation(); pres.slide_width = Inches(13.33); pres.slide_height = Inches(7.5)
    preview_pdf = fitz.open() if reading_preview else None
    font = APP/'assets/fonts/NotoSansSC-Regular.ttf'
    for index, data in enumerate(slides, 1):
        slide = pres.slides.add_slide(pres.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string('F7FAFB')
        image = None
        if data.get('screenshot'):
            image = data['screenshot'].read_bytes()
        elif data.get('pdf'):
            with fitz.open(data['pdf']) as source:
                if not source.page_count: raise ValueError('EMPTY_REPORT_PDF')
                image = source[0].get_pixmap(matrix=fitz.Matrix(1.2,1.2)).tobytes('png')
        body_width = 7.1 if image else 12
        boxes = [(0.7,0.5,12,1.05,data['title'],30,'174C5F'),
                 (0.7,1.85,body_width,4.9,'\n'.join(data['lines']),18 if image else 21,'243C48'),
                 (0.7,7,12,0.3,f'药衡智析 · 运行证据与验收边界 / {index}',11,'567580')]
        for x,y,w,h,text,size,color in boxes:
            box = slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
            frame = box.text_frame; frame.word_wrap = True
            for i,line in enumerate(text.splitlines()):
                para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                para.text = line; para.font.name = 'Noto Sans SC'; para.font.size = Pt(size)
                para.font.color.rgb = RGBColor.from_string(color); para.space_after = Pt(13 if size >= 18 else 0)
        if image:
            slide.shapes.add_picture(BytesIO(image), Inches(8.65), Inches(1.85), height=Inches(4.85))
        slide.notes_slide.notes_text_frame.text = json.dumps(data.get('notes') or {}, ensure_ascii=False, indent=2, default=str)
        if preview_pdf is not None:
            page = preview_pdf.new_page(width=960,height=540)
            page.draw_rect(page.rect,color=None,fill=(.97,.98,.985)); page.insert_font(fontname='Noto',fontfile=str(font))
            title_space = fitz.Rect(50,30,910,100)
            if page.insert_textbox(title_space,data['title'],fontname='Noto',fontsize=26,color=(.09,.30,.37)) < 0:
                raise ValueError('READING_PREVIEW_TITLE_OVERFLOW')
            if page.insert_textbox(fitz.Rect(50,130,575 if image else 910,480),'\n\n'.join(data['lines']),fontname='Noto',fontsize=15 if image else 17,color=(.14,.24,.28)) < 0:
                raise ValueError('READING_PREVIEW_BODY_OVERFLOW')
            if image: page.insert_image(fitz.Rect(620,130,915,480),stream=image,keep_proportion=True)
            page.insert_text((50,515),f'阅读预览 · 真人审核待完成 / {index}',fontname='Noto',fontsize=11)
    path = output_dir/'药衡智析_演示与答辩稿.pptx'; pres.save(path)
    if len(Presentation(path).slides) != 10: raise ValueError('SLIDE_COUNT_MUST_BE_TEN')
    if preview_pdf is not None:
        preview = output_dir/'_reading_preview'; preview.mkdir(exist_ok=True)
        preview_pdf.save(preview/'药衡智析_演示与答辩稿_阅读版.pdf')
        for i,page in enumerate(preview_pdf): page.get_pixmap(matrix=fitz.Matrix(.8,.8)).save(preview/f'slide-{i+1}.png')
        preview_pdf.close()
    return path


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--output-dir',type=Path,required=True)
    parser.add_argument('--manifest',type=Path)
    parser.add_argument('--reports-dir',type=Path)
    parser.add_argument('--reading-preview',action='store_true')
    args = parser.parse_args()
    manifest, loaded = load_run(args.manifest,args.reports_dir)
    path = write_deck(build_slides(manifest,loaded),args.output_dir,args.reading_preview)
    print('Verified PPTX slides: 10; run:',manifest['run_id'],'; reading preview:',args.reading_preview,'; native rendering is separate')
    print(path)

if __name__ == '__main__': main()
