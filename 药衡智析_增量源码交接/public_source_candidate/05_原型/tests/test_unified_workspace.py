"""Real multi-upload regressions; no external model, demo data or runtime writes."""
import importlib
import io
import json
from decimal import Decimal

import pytest

from pharma import data_import, industry, import_pipeline


@pytest.fixture()
def runtime(tmp_path, monkeypatch):
    monkeypatch.setenv('PHARMA_RUNTIME_DIR', str(tmp_path))
    from pharma import config
    importlib.reload(config)
    importlib.reload(data_import)
    importlib.reload(import_pipeline)
    monkeypatch.setattr(industry, 'ENTERPRISE_REGISTRY', tmp_path / 'enterprise_registry.json')
    monkeypatch.setattr(import_pipeline, '_extraction_mapping', lambda h, s: (None, 'deterministic fixture'))
    monkeypatch.setattr(import_pipeline, '_analysis_hypotheses', lambda a: None)
    return tmp_path


def upload(name, rows, data_type='cost_summary', header='工厂,产品,月份,产量,直接材料'):
    record=data_import.create_upload('business', name, (header+'\n'+rows+'\n').encode(), data_type)
    mapping={h:v for h,v in data_import.suggest_mapping(record['meta']['preview']['headers']).items()
             if not h.startswith('_')}
    return record,mapping,{'scenario':'budget' if data_type=='budget' else 'actual'}


def publish(entry, **kwargs):
    return data_import.publish_workspace([entry], **kwargs)


def snapshot(cid, month='2026-02'):
    return industry.analyze_reference(cid, factory='F', product='P', month=month)


def test_successive_uploads_form_one_workspace_with_history_and_budget(runtime):
    january=upload('一月.csv','F,P,2026-01,10,80')
    first=publish(january,enterprise_name='企业A',quantity_unit='件')
    frozen_path=industry._registered()[first['context_id']]
    frozen_bytes=open(frozen_path,'rb').read()
    first_version=industry.resolve_context(first['context_id']).data_snapshot
    february=upload('二月.csv','F,P,2026-02,10,100')
    second=publish(february)
    budget=upload('预算.csv','F,P,2026-02,10,90','budget')
    third=publish(budget)
    assert first['context_id']==second['context_id']==third['context_id']
    assert open(frozen_path,'rb').read()==frozen_bytes  # prior version remains recoverable
    assert first_version != industry.resolve_context(third['context_id']).data_snapshot
    result=snapshot(third['context_id'])
    assert Decimal(result['metrics']['unit_cost']['value'])==Decimal('10')
    assert Decimal(result['metrics']['mom']['value'])==Decimal('25')
    assert Decimal(result['metrics']['budget']['value']) > Decimal('11.1')
    assert len(result['trend'])==2
    state=data_import.workspace_state()
    assert state['status']=='READY' and state['all_files_included']
    assert len(state['imported_files'])==3
    assert state['data_snapshot']==industry.resolve_context(third['context_id']).data_snapshot
    assert industry.context_catalog()['default_context_id']==third['context_id']
    assert '一月.csv:2:直接材料' in str(result['metrics']['mom']['sources'])


def test_complementary_files_and_identical_export_do_not_multiply_quantity_or_amount(runtime):
    material=upload('材料.csv','F,P,2026-02,10,60')
    first=publish(material)
    labor=upload('人工.csv','F,P,2026-02,10,40',header='工厂,产品,月份,产量,直接人工')
    publish(labor)
    # Different source bytes with exactly the same business fact are one fact.
    duplicate=upload('材料重导.csv','F,P,2026-02,10.0,60.00')
    publish(duplicate)
    result=snapshot(first['context_id'])
    assert Decimal(result['metrics']['quantity']['value'])==10
    assert Decimal(result['metrics']['total_cost']['value'])==100
    assert Decimal(result['metrics']['unit_cost']['value'])==10
    assert '材料.csv:2:直接材料' in str(result['metrics']['material']['row_keys'])
    assert '材料重导.csv:2:直接材料' in str(result['metrics']['material']['row_keys'])


@pytest.mark.parametrize('rows,expected',[
    ('F,P,2026-02,10,101','IMPORT_AMOUNT_CONFLICT'),
    ('F,P,2026-02,11,100','IMPORT_QUANTITY_CONFLICT'),
])
def test_conflicts_preserve_last_complete_snapshot_and_are_visible(runtime,rows,expected):
    first=publish(upload('正确.csv','F,P,2026-02,10,100'))
    registry_before=industry.ENTERPRISE_REGISTRY.read_bytes()
    snapshot_before=industry.resolve_context(first['context_id']).data_snapshot
    wrong=upload('冲突.csv',rows)
    with pytest.raises(ValueError,match=expected) as error:publish(wrong)
    data_import.mark_import_status(wrong[0],'PARSE_FAILED',{'parse_error':str(error.value)})
    assert industry.ENTERPRISE_REGISTRY.read_bytes()==registry_before
    assert industry.resolve_context(first['context_id']).data_snapshot==snapshot_before
    state=data_import.workspace_state()
    assert state['status']=='BLOCKED' and not state['all_files_included']
    assert state['context_id']==first['context_id'] and state['issues'][0]['filename']=='冲突.csv'
    assert state['source']=='uploaded'


def test_pending_upload_never_displays_competition_fallback(runtime,monkeypatch):
    monkeypatch.setattr(industry,'_competition_available',lambda:True)
    assert data_import.workspace_state()['context_id']=='pharmaceutical:competition'
    upload('待解析.csv','F,P,2026-02,10,100')
    state=data_import.workspace_state()
    assert state['status']=='PENDING' and state['context_id'] is None
    assert industry.context_catalog()['default_context_id'] is None


def test_scope_and_unit_conflicts_are_not_silently_mixed(runtime):
    publish(upload('A.csv','F,P,2026-01,10,100'),enterprise_name='A',quantity_unit='件')
    new=upload('B.csv','F,P,2026-02,10,100')
    with pytest.raises(ValueError,match='WORKSPACE_ENTERPRISE_CONFLICT'):
        publish(new,enterprise_name='B')
    with pytest.raises(ValueError,match='WORKSPACE_UNIT_CONFLICT'):
        publish(new,quantity_unit='盒')
    hidden_unit=upload('带单位.csv','F,P,2026-02,10,100',header='工厂,产品,月份,产量(盒),直接材料')
    with pytest.raises(ValueError,match='WORKSPACE_UNIT_CONFLICT'):publish(hidden_unit)


def test_old_published_import_migrates_without_losing_knowledge_context(runtime):
    entry=upload('旧版.csv','F,P,2026-01,10,80')
    old=data_import.publish_business(*entry,'企业A','generic_manufacturing','件')
    record=data_import.get_import(entry[0]['id'])
    data_import._save(record,{**record['meta'],'published':old})
    state=data_import.workspace_state()
    assert state['status']=='READY' and state['context_id']==old['context_id']
    assert state['legacy_context_ids']==[old['context_id']]
    publish(upload('新月份.csv','F,P,2026-02,10,100'))
    assert Decimal(snapshot(old['context_id'])['metrics']['mom']['value'])==25


def test_pipeline_keeps_previous_accepted_sources(runtime):
    from pharma.jobs import JobStore
    store=JobStore(runtime/'jobs.sqlite3')
    for month,amount in [('01','80'),('02','100')]:
        entry=upload(month+'.csv',f'F,P,2026-{month},10,{amount}')
        job=store.enqueue('data_parse',{'import_ids':[entry[0]['id']]})
        import_pipeline.run_data_parse(store,store.get(job['id']))
        assert store.get(job['id'])['status']=='SUCCEEDED'
    state=data_import.workspace_state()
    assert len(state['imported_files'])==2 and state['all_files_included']
    assert Decimal(snapshot(state['context_id'])['metrics']['mom']['value'])==25
    assert all(record['meta'].get('business_contract') for record in data_import._business_records())


def test_more_than_ui_page_limit_are_all_queued_and_reported(runtime):
    for number in range(201):
        upload(f'{number}.csv',f'F,P{number},2026-02,10,100')
    assert len(data_import.waiting_imports('business'))==201
    assert len(data_import.workspace_state()['pending_files'])==201


def test_identical_actual_budget_bytes_keep_distinct_roles_and_missing_amount_not_zero(runtime):
    entry=upload('actual.csv','F,P,2026-02,10,100')
    budget=upload('budget.csv','F,P,2026-02,10,100','budget')
    assert budget[0]['id']!=entry[0]['id']
    missing=upload('缺失.csv','F,P,2026-03,10,')
    with pytest.raises(ValueError,match='IMPORT_VALIDATION_FAILED'):
        data_import.publish_workspace([entry,missing])
    assert data_import.workspace_state()['context_id'] is None


def test_missing_dimension_row_is_rejected_instead_of_disappearing(runtime):
    entry=upload('缺少产品.csv','F,P,2026-02,10,100\nF,,2026-03,10,120')
    validation=data_import.validate_business(*entry)
    assert validation['status']=='INVALID'
    assert any(error['row']==3 and '产品' in error['reason'] for error in validation['errors'])
    with pytest.raises(ValueError,match='IMPORT_VALIDATION_FAILED'):publish(entry)


def test_budget_only_waits_for_actual_instead_of_publishing_empty_analysis(runtime):
    budget=upload('只有预算.csv','F,P,2026-02,10,100','budget')
    with pytest.raises(ValueError,match='WORKSPACE_ACTUAL_DATA_REQUIRED'):publish(budget)
    assert data_import.workspace_state()['context_id'] is None
    actual=upload('实际.csv','F,P,2026-02,10,110')
    result=data_import.publish_workspace([budget,actual])
    assert Decimal(snapshot(result['context_id'])['metrics']['budget']['value'])==10


def test_explicit_file_replacement_keeps_original_and_previous_version(runtime):
    original=upload('错误原件.csv','F,P,2026-02,10,100')
    first=publish(original)
    source=data_import.original_path(original[0])
    original_bytes=source.read_bytes()
    old_version=industry.resolve_context(first['context_id']).data_snapshot
    correction=upload('修正版.csv','F,P,2026-02,10,120')
    with pytest.raises(ValueError,match='IMPORT_AMOUNT_CONFLICT'):publish(correction)
    correction[2]['replace_import_id']=original[0]['id']
    updated=publish(correction)
    assert updated['context_id']==first['context_id']
    assert Decimal(snapshot(first['context_id'])['metrics']['unit_cost']['value'])==12
    assert source.read_bytes()==original_bytes
    assert industry.resolve_context(first['context_id']).data_snapshot!=old_version
    state=data_import.workspace_state()
    assert state['status']=='READY' and state['all_files_included']
    assert len(state['imported_files'])==1 and len(state['superseded_files'])==1
    assert state['superseded_files'][0]['replaced_by']==correction[0]['id']
    publish(correction)  # same action can be retried without another retirement
    assert len(data_import.workspace_state()['superseded_files'])==1
    assert data_import.waiting_imports('business')==[]


def test_failed_replacement_retains_old_pointer(runtime):
    original=upload('原件.csv','F,P,2026-02,10,100')
    first=publish(original)
    before=industry.ENTERPRISE_REGISTRY.read_bytes()
    invalid=upload('缺量修订.csv','F,P,2026-02,,120')
    invalid[2]['replace_import_id']=original[0]['id']
    with pytest.raises(ValueError):publish(invalid)
    assert industry.ENTERPRISE_REGISTRY.read_bytes()==before
    assert Decimal(snapshot(first['context_id'])['metrics']['unit_cost']['value'])==10


def test_identical_detail_rows_rejected_without_transaction_ids(runtime):
    detail=upload('重复明细.csv','F,P,2026-02,10,60\nF,P,2026-02,10,60','material_detail')
    with pytest.raises(ValueError,match='完全重复'):publish(detail)


def test_declared_quantity_unit_is_inferred_and_currency_cannot_be_silently_relabeled(runtime):
    entry=upload('盒单位.csv','F,P,2026-02,10,3.3',header='工厂,产品,月份,产量(盒),直接材料(元/盒)')
    result=publish(entry)
    assert result['quantity_unit']=='盒'
    assert Decimal(snapshot(result['context_id'])['metrics']['unit_cost']['value'])==Decimal('3.3')
    foreign=upload('外币.csv','F,P,2026-03,10,100')
    foreign[2]['currency']='USD'
    with pytest.raises(ValueError,match='WORKSPACE_CURRENCY_NOT_SUPPORTED'):publish(foreign)


def test_publish_during_analysis_cannot_mix_version_labels_with_new_numbers(runtime,monkeypatch):
    first=publish(upload('旧.csv','F,P,2026-01,10,100'))
    frozen=industry.resolve_context(first['context_id'])
    publish(upload('新.csv','F,P,2026-02,10,120'))
    monkeypatch.setattr(industry,'resolve_context',lambda _cid:frozen)
    with pytest.raises(ValueError,match='DATA_SNAPSHOT_CHANGED'):snapshot(first['context_id'])


def test_api_default_analysis_uses_all_uploads_and_pending_file_blocks_generation(runtime,monkeypatch):
    from fastapi.testclient import TestClient
    from pharma import api
    from pharma.jobs import JobStore
    monkeypatch.setattr(api,'store',JobStore(runtime/'api-jobs.sqlite3'))
    monkeypatch.setenv('PHARMA_AUTO_EXPLAIN','0')
    monkeypatch.delenv('PHARMA_API_TOKEN',raising=False)
    with TestClient(api.app) as client:
        for month,amount in [('01','80'),('02','100')]:
            record,mapping,options=upload(month+'.csv',f'F,P,2026-{month},10,{amount}')
            response=client.post(f'/api/imports/{record["id"]}/publish',json={'mapping':mapping,'options':options})
            assert response.status_code==201,response.text
        state=client.get('/api/workspace').json()
        assert state['status']=='READY' and len(state['imported_files'])==2
        response=client.post('/api/analyses',json={'month':'2026-02'})
        assert response.status_code==200,response.text
        assert Decimal(response.json()['metrics']['mom']['value'])==25
        upload('待补齐.csv','F,P,2026-03,10,120')
        blocked=client.post('/api/analyses',json={'month':'2026-02'})
        assert blocked.status_code==422 and '尚未' in blocked.json()['error']['message']
        pending=client.get('/api/workspace').json()
        assert pending['context_id']==state['context_id'] and pending['source']=='uploaded'


def test_api_parse_replacement_map_rebuilds_complete_workspace(runtime,monkeypatch):
    from fastapi.testclient import TestClient
    from pharma import api
    from pharma.jobs import JobStore
    store=JobStore(runtime/'api-replace.sqlite3')
    monkeypatch.setattr(api,'store',store)
    monkeypatch.delenv('PHARMA_API_TOKEN',raising=False)
    original=upload('旧成本.csv','F,P,2026-02,10,100')
    publish(original)
    revised=upload('正确成本.csv','F,P,2026-02,10,120')
    with TestClient(api.app) as client:
        submitted=client.post('/api/data/parse',json={'replacements':{revised[0]['id']:original[0]['id']}})
        assert submitted.status_code==202,submitted.text
        job=store.get(submitted.json()['job_id'])
        assert job['input']['import_ids']==[revised[0]['id']]
        import_pipeline.run_data_parse(store,job)
        assert store.get(job['id'])['status']=='SUCCEEDED'
        state=client.get('/api/workspace').json()
        assert state['status']=='READY' and len(state['superseded_files'])==1
        assert Decimal(snapshot(state['context_id'])['metrics']['unit_cost']['value'])==12


def _workbook(second_rows=None, empty_first=False):
    from openpyxl import Workbook
    book=Workbook()
    if empty_first:
        book.active.title='空白封面'
        data=book.create_sheet('成本表')
    else:data=book.active
    data.append(['工厂','产品','月份','产量','直接材料'])
    data.append(['F','P','2026-02',10,100])
    extra=book.create_sheet('另一工作表')
    for row in second_rows or []:extra.append(row)
    stream=io.BytesIO();book.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize('rows',[
    [['工厂','产品','月份','产量','直接材料'],['F','P','2026-03',10,120]],
    [['=SUM(1,2)']],
])
def test_multiple_populated_worksheets_are_retained_but_never_partially_published(runtime,rows):
    payload=_workbook(rows)
    record=data_import.create_upload('business','多表.xlsx',payload)
    assert '多个非空工作表' in record['meta']['preview']['warning']
    mapping=data_import.suggest_mapping(record['meta']['preview']['headers'])
    with pytest.raises(ValueError,match='尚未接入任何工作表'):
        publish((record,mapping,{}))
    assert data_import.original_path(record).read_bytes()==payload
    assert data_import.workspace_state()['context_id'] is None


def test_blank_extra_worksheets_are_ignored_and_first_nonempty_table_is_used(runtime):
    record=data_import.create_upload('business','空白页加成本表.xlsx',_workbook([['  ',None]],empty_first=True))
    assert record['meta']['sheet']=='成本表'
    mapping=data_import.suggest_mapping(record['meta']['preview']['headers'])
    result=publish((record,mapping,{}))
    assert Decimal(snapshot(result['context_id'])['metrics']['unit_cost']['value'])==10
