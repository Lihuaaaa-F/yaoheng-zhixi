import csv, hashlib, json, re, sys
from collections import Counter
from datetime import datetime,timezone
from decimal import Decimal as D
from fractions import Fraction as F
from pathlib import Path
import pymupdf
from docx import Document
ROOT=Path(__file__).resolve().parents[3]
index_path=ROOT/sys.argv[1]
rows=[]; sources=[]
for p in sorted((ROOT/'01_数据/00_原始').rglob('*.csv')):
    sources.append({'path':str(p.relative_to(ROOT)),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
    with p.open(encoding='utf-8-sig',newline='') as f:
        kind='budget' if '预算数据' in p.name else 'cost' if '成本汇总' in p.name else 'materials' if '原材料消耗' in p.name else 'labor' if '人工工时' in p.name else 'expenses' if '制造费用明细' in p.name else None
        if kind: rows.extend(dict(r,_kind=kind) for r in csv.DictReader(f))
def selected(factory,product,period,kind='cost'):
    return [r for r in rows if r['_kind']==kind and r['工厂']==factory and r['产品名称']==product and r['月份'] in period]
def aggregate(factory,product,period,kind='cost'):
    rr=selected(factory,product,period,kind); prefix='预算' if kind=='budget' else ''
    if set(r['月份'] for r in rr)!=set(period):return None
    q=sum(F(r[prefix+'产量(盒)']) for r in rr); total=sum(F(r[prefix+'总成本(元)']) for r in rr)
    return {'quantity':q,'total_cost':total,'unit_cost':total/q,**{k:sum(F(r[prefix+n+'(元/盒)'])*F(r[prefix+'产量(盒)']) for r in rr)/q for k,n in [('materials','直接材料'),('labor','直接人工'),('overhead','制造费用')]}}
def percent(a,b):return (a-b)/b*100 if b else None
def rounded(v,n=2):return f'{D(v.numerator)/D(v.denominator):.{n}f}'
def clean(t):return re.sub(r'\s+','',t)
checks=[]; report_results=[]
def check(case,field,passed,actual=None,expected=None):
    checks.append({'scenario':case,'field':field,'status':'PASS' if passed else 'FAIL','actual':actual,'expected':str(expected) if expected is not None else None})
def numeric(case,field,text,expected):
    if expected is None:
        check(case,field,'N/A' in text or '—' in text,text,None);return
    try:
        s=text.strip().replace('%','').replace(',',''); number=F(s)
        places=len(s.split('.')[1]) if '.' in s else 0
        valid=abs(number-expected)<=F(1,2*10**places)
    except Exception:valid=False
    check(case,field,valid,text,expected)
for entry in json.loads(index_path.read_text()):
    case=entry['scenario']; p=ROOT/'07_交付/业务报告'/entry['job_id']; r=json.loads((p/'record.json').read_text()); audit=json.loads((p/'machine_audit.json').read_text())
    inp=entry['input']; factory,product,month=inp['factory'],inp['product'],inp['month']; quarterly=inp['analysis_type']=='quarterly'
    end=int(month[-2:]);period=[f'{month[:4]}-{i:02}' for i in range(end-2,end+1)] if quarterly else [month]
    offset=3 if quarterly else 1; mom=[]
    for m in period:
        n=int(m[:4])*12+int(m[-2:])-1-offset;mom.append(f'{n//12:04}-{n%12+1:02}')
    yoy=[f'{int(m[:4])-1}-{m[-2:]}' for m in period]
    now=aggregate(factory,product,period); prev=aggregate(factory,product,mom); last=aggregate(factory,product,yoy); budget=aggregate(factory,product,period,'budget')
    d=Document(p/'report.docx'); pdf=pymupdf.open(p/'report.pdf'); pdfpages=[clean(x.get_text()) for x in pdf]; pdftext=''.join(pdfpages); pdfbody=''.join(clean(re.sub(r'^第\s*\d+\s*页\s*共\s*\d+\s*页\s*', '', x.get_text())) for x in pdf)
    docxtext=clean('\n'.join(x.text for x in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for rr in t.rows for c in rr.cells))
    for kind,filename in [('docx','report.docx'),('pdf','report.pdf'),('audit','machine_audit.json')]:
        digest=hashlib.sha256((p/filename).read_bytes()).hexdigest()
        check(case,kind+'.index_hash',digest==entry['artifacts'][kind]['sha256'],digest,entry['artifacts'][kind]['sha256'])
        check(case,kind+'.record_hash',digest==r[kind]['sha256'],digest,r[kind]['sha256'])
    check(case,'snapshot.record_audit_identity',r['snapshot']==audit['snapshot'])
    check(case,'bindings.index_record_audit_identity',entry['artifacts']['docx']['bindings']==r['docx']['bindings']==audit['bindings'])
    for key in ('quantity','unit_cost','total_cost','materials','labor','overhead'):
        actual=r['snapshot']['metrics'][key]['value']; check(case,'snapshot.'+key,abs(F(actual)-now[key])<F(1,10**12),actual,now[key])
    for i,key in enumerate(('quantity','unit_cost','total_cost','materials','labor','overhead'),1):
        expected=[now[key],prev[key],percent(now[key],prev[key]),last[key],percent(now[key],last[key]),budget[key],percent(now[key],budget[key])]
        cells=d.tables[1].rows[i].cells
        for j,value in enumerate(expected,1):numeric(case,f'overview.{key}.column{j}',cells[j].text,value)
        # Read order is normalized for whitespace only, not rearranged or reconstructed.
        rowtext=clean(''.join(c.text for c in cells))
        check(case,f'pdf.overview.{key}.same_row',rowtext in pdftext,rowtext)
    for i,key in enumerate(('materials','labor','overhead'),1):
        delta=now[key]-prev[key]; share=now[key]/now['unit_cost']*100; contribution=delta/(now['unit_cost']-prev['unit_cost'])*100
        for j,v in enumerate((now[key],share,delta,contribution),1):numeric(case,f'structure.{key}.column{j}',d.tables[2].rows[i].cells[j].text,v)
        for label,base in [('mom',prev),('yoy',last),('budget',budget)]:
            bound=r['snapshot']['elements'][i-1]['comparisons'][label]['unit']; numerator=now[key]-base[key]; denominator=now['unit_cost']-base['unit_cost']
            for f,v in [('numerator',numerator),('denominator',denominator),('contribution',numerator/denominator*100 if denominator else None)]:
                actual=bound[f];check(case,f'bound.{key}.{label}.{f}',(actual is None if v is None else abs(F(actual)-v)<F(1,10**12)),actual,v)
    # Every current-period raw-material detail row, including all quarterly months.
    material_rows=selected(factory,product,period,'materials')
    actual_material={(rr.cells[0].text,rr.cells[1].text):rr for rr in d.tables[3].rows[1:]}
    check(case,'material_detail.row_count',len(material_rows)==len(actual_material),len(actual_material),len(material_rows))
    for mr in material_rows:
        rr=actual_material.get((mr['月份'],mr['原材料名称']))
        if rr is None:check(case,'material_detail.missing',False,mr['原材料名称']);continue
        numeric(case,'material_detail.'+mr['月份']+'.'+mr['原材料名称']+'.unit',rr.cells[2].text,F(mr['单位消耗成本(元/盒)']))
        numeric(case,'material_detail.'+mr['月份']+'.'+mr['原材料名称']+'.total',rr.cells[3].text,F(mr['原材料总成本(元)']))
    for item in r['snapshot']['materials_summary']:
        independent=[]
        for pp in (period,mom):
            mr=[x for x in selected(factory,product,pp,'materials') if x['原材料名称']==item['name']]
            independent.append(sum(F(x['原材料总成本(元)']) for x in mr)/sum(F(x['产量(盒)']) for x in mr))
        a,b=independent; share=(a-b)/(now['materials']-prev['materials'])*100
        for key,value in [('current',a),('previous',b),('delta',a-b),('contribution',share)]:
            check(case,'material_driver.'+item['name']+'.'+key,abs(F(item[key])-value)<F(1,10**12),item[key],value)
        if item in r['snapshot']['materials_summary'][:4]:
            phrase='占材料增量'+rounded(share)+'%'
            check(case,'docx.material_driver.'+item['name']+'.contribution',phrase in docxtext,phrase)
            check(case,'pdf.material_driver.'+item['name']+'.contribution',phrase in pdfbody,phrase)
    # Independently weighted labor measures; average hourly wage is a conversion.
    labor_values=[]
    for pp in (period,mom):
        lr=selected(factory,product,pp,'labor'); hours=sum(F(x['总工时(小时)']) for x in lr);q=sum(F(x['产量(盒)']) for x in lr);wages=sum(F(x['直接人工总额(元)']) for x in lr);days=sum(F(x['生产人数(人)'])*F(x['工作天数(天)']) for x in lr)
        labor_values.append([wages/q,hours/q*10000,wages/hours,q/days])
    for i,(a,b) in enumerate(zip(*labor_values),1):
        for j,v in enumerate((a,b,percent(a,b)),1):numeric(case,f'labor.row{i}.column{j}',d.tables[4].rows[i].cells[j].text,v)
    left=aggregate('中药二厂',product,period); right=aggregate('中药一厂',product,period); overall=left['unit_cost']-right['unit_cost']
    for i,key in enumerate(('unit_cost','materials','labor','overhead'),1):
        delta=left[key]-right[key]
        for j,v in enumerate((left[key],right[key],delta,percent(left[key],right[key])),1):numeric(case,f'benchmark.{key}.column{j}',d.tables[8].rows[i].cells[j].text,v)
        if key!='unit_cost':
            expected=rounded(delta/overall*100)+'%'
            check(case,f'benchmark.{key}.contribution_docx',expected in docxtext,expected)
            check(case,f'benchmark.{key}.contribution_pdf',expected in pdftext,expected)
    bridge={'quantity_effect':(now['quantity']-budget['quantity'])*budget['unit_cost'],'unit_cost_effect':now['quantity']*(now['unit_cost']-budget['unit_cost']),'total_delta':now['total_cost']-budget['total_cost']}
    for k,v in bridge.items():
        actual=r['snapshot']['budget_bridge'][k];check(case,'budget_bridge.'+k,abs(F(actual)-v)<F(1,10**12),actual,v)
        check(case,'docx.budget_bridge.'+k,rounded(v) in docxtext,rounded(v))
        check(case,'pdf.budget_bridge.'+k,rounded(v) in pdftext,rounded(v))
    actual_months=[]
    for rr in d.tables[6].rows[1:]:
        m=rr.cells[0].text; actual_months.append(m); expected=aggregate(factory,product,[m])
        check(case,'trend.no_future.'+m,m<=month,m,month)
        for j,k in enumerate(('unit_cost','quantity','total_cost'),1):numeric(case,f'trend.{m}.{k}',rr.cells[j].text,expected[k])
        check(case,'pdf.trend.'+m,clean(''.join(c.text for c in rr.cells)) in pdftext)
    expected_months=[f'{month[:4]}-{i:02}' for i in range(max(1,end-5),end+1)]
    check(case,'trend.available_months',actual_months==expected_months,actual_months,expected_months)
    # Confirm business reader exports contain their computed table rows, not only registered values.
    for ti in (2,3,4,5,8):
        for ri,rr in enumerate(d.tables[ti].rows[1:],1):
            rowtext=clean(''.join(c.text for c in rr.cells))
            check(case,f'pdf.table{ti}.row{ri}.same_as_docx',rowtext in pdftext,rowtext)
    check(case,'degradation_disclosed_docx','本次采用基础分析，原因解释待复核。' in docxtext)
    check(case,'degradation_disclosed_pdf','本次采用基础分析，原因解释待复核。' in pdftext)
    check(case,'no_report_approval_claim',entry['acceptance']['overall']!='PASS' and r['acceptance']['overall']!='PASS')
    report_results.append({'scenario':case,'job_id':entry['job_id'],'product':product,'period':period,'pdf_pages':len(pdf),'model_live':entry['model_live'],'report_acceptance':entry['acceptance']['overall']})
    pdf.close()
result={'created_at':datetime.now(timezone.utc).isoformat(),'index':str(index_path.relative_to(ROOT)),'index_sha256':hashlib.sha256(index_path.read_bytes()).hexdigest(),'method':'Read-only original CSV Fraction arithmetic; DOCX table cells checked against source arithmetic at displayed precision; whitespace-only PDF table reading-order matching; narrative matching additionally removes explicit page-number lines; independent SHA-256 against artifact index and record. No production calculation functions used.','scope':'Four final report files, period comparisons, material/labor detail, cross-factory direction and contribution, budget bridge, trend, and quarter weighting. Human attribution, readability and visual grading excluded.','sources':sources,'reports':report_results,'checks':checks,'check_count':len(checks),'failure_count':sum(x['status']=='FAIL' for x in checks),'human_attribution_score':'PENDING','human_readability':'PENDING','visual_quality':'PENDING_PARENT_REVIEW'}
result['numeric_file_audit_status']='PASS' if not result['failure_count'] else 'FAIL'
(ROOT/'06_评测/incremental_20260916/data/report_independent_audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in result.items() if k not in ('sources','checks')},ensure_ascii=False,indent=2))
for failure in [x for x in checks if x['status']=='FAIL'][:30]:print(json.dumps(failure,ensure_ascii=False))
