"""Isolated synthetic upload -> API queue -> worker -> real DOCX/PDF downloads.

Requires caller-provided PHARMA_RUNTIME_DIR / PHARMA_ARTIFACTS_DIR. No model key,
no human scoring, and no external notifications. Invoked by the opt-in E2E test.
"""
import hashlib
import io
import json
import os
import sqlite3
from pathlib import Path
from decimal import Decimal


def main():
    assert os.environ.get('PHARMA_RUNTIME_DIR') and os.environ.get('PHARMA_ARTIFACTS_DIR')
    from fastapi.testclient import TestClient
    from docx import Document
    import fitz
    from pharma import api, worker
    client = TestClient(api.app)
    csv = '工厂,产品名称,月份,产量(盒),直接材料(元/盒),直接人工(元/盒),制造费用(元/盒)\n'
    csv += ''.join(f'候选测试厂,独立测试制剂,2031-{m:02d},100,10,2,3\n' for m in range(1,7))
    uploaded = client.post('/api/imports/uploads', data={'kind':'business','data_type':'cost_summary'},
        files={'file':('候选测试件.csv',csv.encode(),'text/csv')})
    assert uploaded.status_code == 201, uploaded.text
    parsed = client.post('/api/data/parse',json={'enterprise_name':'独立测试企业','industry_id':'pharmaceutical','quantity_unit':'盒'})
    assert parsed.status_code == 202, parsed.text
    parse_id=parsed.json()['job_id']; worker.process_job(api.store,api.store.get(parse_id))
    outcome=client.get('/api/jobs/'+parse_id).json()
    assert outcome['status']=='SUCCEEDED', outcome
    cid=outcome['result']['published']['context_id']
    results=[]
    for kind in ('monthly','quarterly','special'):
        selection={'context_id':cid,'factory':'候选测试厂','product':'独立测试制剂','month':'2031-06','analysis_type':kind,'topic':'  '}
        pre=client.post('/api/reports/preflight',json=selection)
        assert pre.json()['ready'], pre.text
        request=client.post('/api/reports',json=selection)
        assert request.status_code==202, request.text
        jid=request.json()['job_id'];worker.process_job(api.store,api.store.get(jid))
        job=client.get('/api/jobs/'+jid).json();result=job['result']
        assert job['status']=='DEGRADED',job # no key or knowledge: honest degradation
        assert result['execution_status']=='COMPLETED' and result['human_review_status']=='PENDING'
        assert result['narrative']['status']=='DEGRADED' and not result['narrative']['model_live']
        assert Decimal(result['snapshot']['metrics']['unit_cost']['value'])==15
        assert Decimal(result['snapshot']['metrics']['total_cost']['value'])==(4500 if kind=='quarterly' else 1500)
        if kind=='special':assert job['input']['topic']=='成本变化与证据核查'
        check=result['docx']['verification'];assert check['status']=='PASS' and check['numeric_bindings_checked']>=70 and not check['residual_placeholders']
        downloads={}
        for fmt in ('docx','pdf'):
            file=client.get('/api/artifacts/'+result[fmt]['artifact_id'])
            assert file.status_code==200 and hashlib.sha256(file.content).hexdigest()==result[fmt]['sha256']
            if fmt=='docx':
                doc=Document(io.BytesIO(file.content));text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
            else:
                pdf=fitz.open(stream=file.content,filetype='pdf');assert len(pdf)>0;text='\n'.join(p.get_text() for p in pdf)
            assert '独立测试制剂' in text and '15.00' in text and '用户导入' in text
            assert '{{' not in text and '[[metric:' not in text
            downloads[fmt]=result[fmt]['sha256']
        results.append({'kind':kind,'job_id':jid,'status':job['status'],'hashes':downloads})
    db=Path(os.environ['PHARMA_RUNTIME_DIR'])/'model_calls.sqlite3'
    # Some versions use model.sqlite3: inspect all ledgers rather than assume name.
    calls=0
    for db in Path(os.environ['PHARMA_RUNTIME_DIR']).glob('*.sqlite3'):
        with sqlite3.connect(db) as conn:
            if conn.execute("select 1 from sqlite_master where name='calls'").fetchone():calls+=conn.execute('select count(*) from calls').fetchone()[0]
    assert calls==0
    print(json.dumps({'status':'PASS','scope':'synthetic import / rules degraded / real artifacts','model_calls':calls,'reports':results},ensure_ascii=False))


if __name__=='__main__':main()
