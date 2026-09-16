import json,re,hashlib,sys
from pathlib import Path
import fitz
from docx import Document
from pharma.config import ROOT,PACKAGE
index=ROOT/(sys.argv[1] if len(sys.argv)>1 else '06_评测/incremental_20260916/scenario_reports_delivery.json')
normalize=lambda value:re.sub(r'\s+','',str(value))
reports=[]
for item in json.loads(index.read_text()):
    folder=ROOT/'07_交付/业务报告'/item['job_id'];rp=folder/'record.json';record=json.loads(rp.read_text())
    snap=record['snapshot'];n=record['narrative'];evidence={e['evidence_id']:e for e in record['evidence']['evidence']}
    with fitz.open(folder/'report.pdf') as pdf:pdfpages=[p.get_text(sort=True) for p in pdf]
    pdftext=normalize(''.join(pdfpages));doc=Document(folder/'report.docx');doctext=normalize('\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for row in t.rows for c in row.cells]))
    checks=[];used=[]
    for finding in n['findings']:
        for ref in finding.get('evidence_refs',[]):
            e=evidence[ref];source=PACKAGE/'03_制药知识文档'/e['source']
            with fitz.open(source) as original:
                page=original[e['page']-1].get_text(sort=True)
                original_text='\n'.join(p.get_text(sort=True) for p in original)
            quote=finding['evidence_quotes'][ref];citation=e['source']+' · '+(e.get('section') or e.get('heading') or '相关章节')+' · '+(e.get('location') or f"第{e['page']}页")
            substantive=finding.get('text','')
            tests={'quote_found_on_original_page':normalize(quote) in normalize(page),'product_matches_declared_section':snap['product'] in e.get('products',[]) or (e.get('scope')=='general' and 'GMP' in e['source']),'factory_no_conflict':not e.get('factory') or e['factory']==snap['factory'],'event_within_period':not e.get('event_period') or snap['period']['start']<=e['event_period']<=snap['period']['end'],'document_version_found_in_original':not e.get('document_version') or normalize('版本:'+e['document_version']) in normalize(original_text),'document_effective_before_report_period':not e.get('effective_date') or e['effective_date'][:7]<=snap['period']['start'],'short_quote':4<=len(quote)<=180,'citation_present_in_docx':normalize(citation) in doctext,'citation_present_in_pdf':normalize(citation) in pdftext,'hypothesis_marked_uncertain':finding['claim_type']!='hypothesis' or (finding.get('hypothesis') is True and '可能' in substantive and bool(finding.get('missing_evidence'))),'no_measured_yield_decline_assertion': '本期实际收率下降' not in substantive or '不能证明本期实际收率下降' in substantive}
            limits=[]
            if not e.get('specification'):limits.append('该引用切片未单独声明包装规格，作为该产品工艺背景；未以其参数替代当期成本计算')
            if not e.get('factory'):limits.append('配方文档未单独声明工厂，作为该产品配方背景；未据此证明本厂实际工艺偏差')
            used.append({'evidence_id':ref,'citation':citation,'quote':quote,'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'document_version':e.get('document_version'),'effective_date':e.get('effective_date'),'scope_limits':limits,'checks':tests,'status':'PASS_WITH_LIMITS' if all(tests.values()) and limits else 'PASS' if all(tests.values()) else 'FAIL'})
    recipe=PACKAGE/'03_制药知识文档'/('产品配方文档_'+snap['product']+'.pdf')
    with fitz.open(recipe) as original:recipe_text=normalize(original[0].get_text(sort=True))
    components={'银黄口服液':['10ml','10支'],'板蓝根颗粒':['10g','20袋'],'六味地黄胶囊':['0.3g','60粒']}[snap['product']]
    specification={'snapshot':snap['specification'],'independent_product_document':recipe.name+' · 第1页','components_match':all(c in recipe_text for c in components),'note':'核对产品配方规格；工艺切片缺包装规格的引用仍标背景限制'}
    tasks=[]
    for f in n['findings']:
        if f.get('suggestion'):
            required=('verification_target','expected_evidence','responsible_role','department','deadline_basis')
            t={'section':f.get('section'),'action':f['suggestion'],'required_fields':{k:bool(f.get(k)) for k in required},'priority_valid':f.get('priority') in ('high','medium','low'),'unknown_assignee_not_fabricated':f.get('responsible_role')=='待分配','action_in_docx':normalize(f['suggestion']) in doctext,'action_in_pdf':normalize(f['suggestion']) in pdftext,'target_in_docx':normalize(f['verification_target']) in doctext,'target_in_pdf':normalize(f['verification_target']) in pdftext}
            t['status']='PASS' if all(t['required_fields'].values()) and all(t[k] for k in ('priority_valid','action_in_docx','action_in_pdf','target_in_docx','target_in_pdf')) else 'FAIL';tasks.append(t)
    retrieved=[{'source':e['source'],'page':e['page'],'products':e.get('products'),'event_period':e.get('event_period')} for e in evidence.values()]
    report={'scenario':item['scenario'],'job_id':item['job_id'],'record_sha256':hashlib.sha256(rp.read_bytes()).hexdigest(),'docx_sha256':hashlib.sha256((folder/'report.docx').read_bytes()).hexdigest(),'pdf_sha256':hashlib.sha256((folder/'report.pdf').read_bytes()).hexdigest(),'model':n['model'],'model_live':n['model_live'],'generation_mode':n['generation_mode'],'used_citations':used,'retrieved_context':retrieved,'specification_crosscheck':specification,'actions':tasks,'human_attribution_score':'PENDING','human_readability':'PENDING','rendered_visual_quality':'NOT_ASSESSED_BY_THIS_TEXT_AND_CITATION_AUDIT'}
    if item['scenario']=='S2':report['specific_regression']={'correct_process_page3_available':any(e['source']=='生产工艺文档_中药一厂.pdf' and e['page']==3 for e in evidence.values()),'silver_process_page2_excluded':not any(e['source']=='生产工艺文档_中药一厂.pdf' and e['page']==2 for e in evidence.values()),'actual_mechanism_citation':'本次采用同产品配方第2页收率与单耗段落；并未冒称引用工艺第3页'}
    if item['scenario']=='S3':
        current_events=[e for e in evidence.values() if e.get('event_period')==snap['month']]
        event_hypotheses=[f for f in n['findings'] if f['claim_type']=='hypothesis' and f.get('section')=='overhead']
        report['specific_regression']={'current_event_available':bool(current_events),'current_event_hypothesis_present':bool(event_hypotheses),'net_decline_not_inferred':all('不等于本月净减产' in f['text'] for f in event_hypotheses),'raw_event_numbers_not_in_body':all(not any(token in f['text'] for token in ('8500','8,500','12000','24')) for f in event_hypotheses)}
        report['coverage_limit']='已形成当期维修事件支持的可能机制，仍缺批次装量、损耗、停工和入账凭证，不能视为已证实因果。' if event_hypotheses else '本次未生成设备原因假设，不能宣称当期维修归因已建立。'
    report['status']='PASS_WITH_LIMITS' if all(c['status']!='FAIL' for c in used) and specification['components_match'] and all(t['status']=='PASS' for t in tasks) else 'FAIL'
    reports.append(report)
out={'scope':'独立重新打开原PDF页、报告DOCX/PDF正文核对真实引用及任务字段；不替代人工归因或视觉评分','index':str(index.relative_to(ROOT)),'reports':reports,'overall':'PASS_WITH_LIMITS' if all(r['status']!='FAIL' for r in reports) else 'FAIL'}
path=ROOT/'06_评测/incremental_20260916/retrieval/report_citation_audit.json';path.write_text(json.dumps(out,ensure_ascii=False,indent=2))
print(json.dumps({'overall':out['overall'],'reports':[(r['scenario'],r['status'],len(r['used_citations']),len(r['actions'])) for r in reports]},ensure_ascii=False))
