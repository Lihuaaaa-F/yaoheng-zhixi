"""图表编号/表格列宽/核心结论回归（2026-09-23 用户三项反馈）。

① 动态表格出"表几-几"题注、图出"图几-几"图题（图下方）；模板原生表
   按用户裁定不加题注（赛题模板固定结构）；
② _fit_table_columns 数值列保足渲染宽度（数字绝不折行）、文本列可压缩；
③ _key_conclusion_lines 结构化核心结论（总量/要素/根因/对标/行动），
   缺数据跳过不编造、不含 None。
"""
import re
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pharma import reports


def _marked_table(doc, marker_index):
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = '药材'
    bm = OxmlElement('w:bookmarkStart')
    bm.set(qn('w:id'), str(40000 + marker_index)); bm.set(qn('w:name'), 'YHU_TBL_' + str(marker_index))
    be = OxmlElement('w:bookmarkEnd'); be.set(qn('w:id'), str(40000 + marker_index))
    cp = t.rows[0].cells[0].paragraphs[0]._p
    cp.insert(0, bm); cp.append(be)
    return t


def test_number_and_caption_marks_figures_and_dynamic_tables_only():
    doc = Document()
    doc.add_paragraph('一、封面与基本信息')
    # 模板原生表（无 YHU_TBL 书签）：不加题注（用户裁定）
    native = doc.add_table(rows=1, cols=2)
    native.rows[0].cells[0].text = '指标'; native.rows[0].cells[1].text = '本月实际'
    # 动态表（书签标记）→ 表1-1 题注
    _marked_table(doc, 0)
    # 图：裸 drawing 元素 + 旧格式图题（编号逻辑只看元素存在）
    pic = doc.add_paragraph(); pic._p.append(OxmlElement('w:drawing'))
    doc.add_paragraph('图｜副题｜趋势图示意')
    reports._number_and_caption(doc, ['趋势图示意'])
    texts = [p.text for p in doc.paragraphs]
    assert '表1-1 趋势图示意' in texts
    assert '图1-1 趋势图示意' in texts
    captioned = [t for t in texts if re.match(r'^表\d+-\d+', t)]
    assert captioned == ['表1-1 趋势图示意']  # 原生表未获题注


def test_fit_table_columns_protects_numeric_width():
    doc = Document()
    t = doc.add_table(rows=4, cols=4)
    for j, h in enumerate(['指标', '本月实际', '环比变动', '同比变动']):
        t.rows[0].cells[j].text = h
    rows = [['其中：直接材料（元/盒）', '614950.00', '-12.50%', '2.69%'],
            ['其中：直接人工（元/盒）', '72800.00', '1.20%', '3.10%'],
            ['单位成本合计（元/盒）', '614950.00', '-12.50%', '2.69%']]
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            t.rows[i].cells[j].text = v
    assert reports._fit_table_columns(t, int(15.9 * 360000))
    grid = t._tbl.find(qn('w:tblGrid'))
    widths = [int(c.get(qn('w:w'))) for c in grid.findall(qn('w:gridCol'))]
    # 数值列宽 ≥ 最长数字渲染需求（9pt：614950.00 = 9×99+140≈1031）
    assert widths[1] >= 980 and widths[2] >= 700 and widths[3] >= 700
    # 文本列不被挤死（≥4 个 CJK 宽）
    assert widths[0] >= 720
    assert sum(widths) <= int(15.9 * 360000 / 635)


def test_key_conclusions_rich_and_none_free():
    snapshot = {'metrics': {'unit_cost': {'value': 17.57}, 'total_cost': {'value': 614950.0}, 'quantity': {'value': 35000}},
                'period_changes': {'unit_cost': {'mom': {'rate': '-2.87'}, 'yoy': {'rate': '2.69'}}},
                'elements': [{'name': '直接材料', 'unit_delta': '-0.40'},
                             {'name': '直接人工', 'unit_delta': '-0.04'},
                             {'name': '制造费用', 'unit_delta': '-0.08'}],
                'attribution': {'status': 'PASS',
                                'ranking': [{'cause': '要素：材料', 'direction': '下降', 'ep_pct': '76.9%', 'label': '中'}],
                                'did': {'parallel_trend': '稳健'}}}
    lines = reports._key_conclusion_lines(snapshot, {'findings': [{'suggestion': '核对采购合同与投料记录'}]})
    assert len(lines) >= 4
    assert any('17.57' in l for l in lines)
    assert any('直接材料' in l for l in lines)
    assert any('76.9%' in l for l in lines)
    assert any('核对采购合同' in l for l in lines)
    assert all('None' not in l for l in lines)
    # 缺数据源时条目优雅跳过，不编造
    lean = reports._key_conclusion_lines({'metrics': {}, 'elements': []}, {'findings': []})
    assert lean == []
