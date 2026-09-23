# -*- coding: utf-8 -*-
"""2026-09-23 全仓审查修复的回归测试（docs/audits/20260923-1f78b2f）。

对应缺陷编号与验收要求（均来自审查报告的"验收"栏）：
- AUD-DATA-01：XLSX 无缓存公式单元格金额不再静默丢失——行级警告、
  金额全空/关键列空值率超阈值判 INVALID（xlsx 反例）。
- AUD-DATA-02：能力预览 available 由数据决定；发布注册失败不留孤儿目录。
- AUD-BENCH-02：benchmark() 缺省方向=本单位−对标厂（与报告路径同向），
  方向断言测试。
- AUD-TPL-01：安装模板同比补绑定按表头/首列文字定位（行序打乱反例）。
- AUD-BENCH-01：generate(allow_model=False) 不发起模型调用即返回规则结果；
  cached_generation 只读缓存不写。
"""
import importlib
import io
import json
from decimal import Decimal

import pytest

from pharma import data_import, import_pipeline, model_settings

WIDE_CSV = ('工厂,产品名称,月份,产量(盒),直接材料(元/盒),直接人工(元/盒),制造费用(元/盒),单位成本(元/盒),总成本(元)\n'
            '中药一厂,银黄口服液,2026-01,10000,3.10,1.00,0.90,5.00,50000\n'
            '中药一厂,银黄口服液,2026-02,10100,3.50,1.00,0.90,5.40,54540\n')


@pytest.fixture
def isolated_runtime(tmp_path, monkeypatch):
    monkeypatch.setenv('PHARMA_RUNTIME_DIR', str(tmp_path))
    monkeypatch.delenv('PHARMA_SHOW_TEST_CONTEXTS', raising=False)
    monkeypatch.setenv('PHARMA_MODEL_KEY_FILE', str(tmp_path / 'no-key-in-tests.key'))
    monkeypatch.delenv('PHARMA_API_KEY', raising=False)
    monkeypatch.delenv('PHARMA_MODEL_ROUTES', raising=False)
    import pharma.config as config
    importlib.reload(config)
    for module in (data_import, model_settings, import_pipeline):
        importlib.reload(module)
    import pharma.industry as industry
    monkeypatch.setattr(industry, 'ENTERPRISE_REGISTRY', tmp_path / 'enterprise_registry.json')
    monkeypatch.setattr(import_pipeline, '_extraction_mapping', lambda h, s: (None, '测试：预设映射'))
    monkeypatch.setattr(import_pipeline, '_analysis_hypotheses', lambda a: None)
    return tmp_path


def _xlsx_payload(amount_formulas: dict[str, bool]) -> bytes:
    """构造成本宽表 xlsx；amount_formulas 指定哪些金额列写公式（无缓存值）。"""
    import openpyxl
    headers = ['工厂', '产品名称', '月份', '产量(盒)', '直接材料(元/盒)', '直接人工(元/盒)', '制造费用(元/盒)']
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    rows = [('中药一厂', '银黄口服液', '2026-01', '10000', '3.10', '1.00', '0.90'),
            ('中药一厂', '银黄口服液', '2026-02', '10100', '3.50', '1.00', '0.90'),
            ('中药一厂', '银黄口服液', '2026-03', '10200', '3.60', '1.05', '0.92')]
    for values in rows:
        out = []
        for header, value in zip(headers, values):
            if amount_formulas.get(header):
                col = openpyxl.utils.get_column_letter(headers.index('产量(盒)') + 1)
                out.append(f'={col}{ws.max_row + 2}*1')  # 公式：读回无缓存值
            else:
                out.append(value)
        ws.append(out)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _mapping_for(record):
    suggested = record['meta']['preview'].get('suggested_mapping') or {}
    return {h: v for h, v in suggested.items() if not str(h).startswith('_')}


# ---------- AUD-BENCH-02：benchmark 缺省方向 ----------

def test_benchmark_default_direction_matches_report_path():
    from pharma.metrics import benchmark, analyze
    analyze('中药一厂', '银黄口服液', '2026-05')  # 确保数据快照已建立（CI 冷缓存）
    comp = benchmark('银黄口服液', '2026-05', factory='中药一厂')
    assert (comp['left'], comp['right']) == ('中药一厂', '中药二厂')
    assert comp['direction'] == '中药一厂−中药二厂，以中药二厂为分母'
    only_left = benchmark('银黄口服液', '2026-05', left='中药一厂')
    assert only_left['right'] == '中药二厂'
    only_right = benchmark('银黄口服液', '2026-05', right='中药二厂')
    assert only_right['left'] == '中药一厂'
    # 显式两厂不受缺省逻辑影响
    explicit = benchmark('银黄口服液', '2026-05', left='中药二厂', right='中药一厂')
    assert explicit['direction'].startswith('中药二厂−中药一厂')


# ---------- AUD-DATA-01：XLSX 公式无缓存值不静默丢失 ----------

def test_xlsx_partial_formula_amounts_get_row_warning_and_invalid(isolated_runtime):
    payload = _xlsx_payload({'直接材料(元/盒)': True})  # 材料列全部为公式
    record = data_import.create_upload('business', '成本汇总.xlsx', payload, 'cost_summary')
    result = data_import.validate_business(record, _mapping_for(record), {})
    # 行级警告：每一行都点名公式列
    assert any('为空（公式无缓存值或漏填）' in w and '直接材料' in w for w in result['warnings'])
    assert any(w.startswith('第 2 行：') for w in result['warnings'])
    # 空值率 3/3 超过 50% 阈值 → INVALID，错误指明原因与出路
    assert result['status'] == 'INVALID'
    assert any('空值率 3/3' in e['reason'] and '公式' in e['reason'] for e in result['errors'])


def test_xlsx_all_formula_amounts_invalid_with_capability_note(isolated_runtime):
    payload = _xlsx_payload({'直接材料(元/盒)': True, '直接人工(元/盒)': True, '制造费用(元/盒)': True})
    record = data_import.create_upload('business', '成本汇总.xlsx', payload, 'cost_summary')
    result = data_import.validate_business(record, _mapping_for(record), {})
    assert result['status'] == 'INVALID'
    assert any('未解析到任何金额数据' in e['reason'] for e in result['errors'])
    caps = {c['capability']: c for c in result['capabilities']}
    assert caps['total_cost_analysis']['available'] is False


def test_few_empty_amount_cells_stay_valid_with_warning_only(isolated_runtime):
    # 少量空单元格（1/3，未过阈值）：VALID 但逐行警告可见，不再静默
    import openpyxl
    headers = ['工厂', '产品名称', '月份', '产量(盒)', '直接材料(元/盒)', '直接人工(元/盒)', '制造费用(元/盒)']
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(headers)
    ws.append(['中药一厂', '银黄口服液', '2026-01', '10000', '3.10', '1.00', '0.90'])
    ws.append(['中药一厂', '银黄口服液', '2026-02', '10100', None, '1.00', '0.90'])
    ws.append(['中药一厂', '银黄口服液', '2026-03', '10200', '3.60', '1.05', '0.92'])
    buf = io.BytesIO(); wb.save(buf)
    record = data_import.create_upload('business', '汇总.xlsx', buf.getvalue(), 'cost_summary')
    result = data_import.validate_business(record, _mapping_for(record), {})
    assert result['status'] == 'VALID'
    assert any('第 3 行：列“直接材料(元/盒)”为空' in w for w in result['warnings'])
    assert result['statistics']['cost_cells'] >= 4  # 其余金额仍参与聚合


# ---------- AUD-DATA-02：能力预览语义 + 发布失败不留孤儿目录 ----------

def test_capability_preview_available_follows_data():
    caps = {c['capability']: c for c in data_import._capabilities({'2026-01'}, {}, {})}
    assert caps['total_cost_analysis']['available'] is False
    amounts = {('f', 'p', '2026-01', 'actual', 'material'): 1}
    caps = {c['capability']: c for c in data_import._capabilities({'2026-01'}, {}, amounts)}
    assert caps['total_cost_analysis']['available'] is True
    # 缺产量：单位成本不可用，且 reason 对齐注册口径（发布需产量）
    assert caps['unit_cost_analysis']['available'] is False
    assert '发布' in caps['unit_cost_analysis']['reason'] and '产量' in caps['unit_cost_analysis']['reason']
    quantities = {('f', 'p', '2026-01', 'actual'): 10}
    caps = {c['capability']: c for c in data_import._capabilities({'2026-01'}, quantities, amounts)}
    assert caps['unit_cost_analysis']['available'] is True


def test_publish_failure_cleans_orphan_enterprise_dir(isolated_runtime, tmp_path, monkeypatch):
    import pharma.industry as industry
    record = data_import.create_upload('business', '汇总.csv', WIDE_CSV.encode('utf-8'), 'cost_summary')
    monkeypatch.setattr(industry, 'register_enterprise',
                        lambda *a, **k: (_ for _ in ()).throw(ValueError('EMPTY_DATASET')))
    with pytest.raises(ValueError, match='EMPTY_DATASET'):
        data_import.publish_business(record, _mapping_for(record), {}, '测试企业', 'pharmaceutical', '盒')
    enterprises = data_import.IMPORTS_ROOT / 'enterprises'
    leftovers = list(enterprises.iterdir()) if enterprises.is_dir() else []
    assert leftovers == [], f'注册失败后残留企业目录：{leftovers}'


# ---------- AUD-TPL-01：同比补绑定按表头文字定位 ----------

def test_install_template_yoy_binding_locates_by_header_text(isolated_runtime, tmp_path, monkeypatch):
    import pharma.reports as reports
    from docx import Document
    monkeypatch.setattr(reports, 'RUNTIME_TEMPLATES', tmp_path / 'templates')
    doc = Document()
    for title in ('一、封面与基本信息', '二、总成本概览', '三、成本要素明细分析',
                  '四、重点产品专项分析', '五、对标分析', '六、总结与建议'):
        doc.add_paragraph(title)
    doc.add_paragraph('{{产品名称}} {{分析月份}}')
    headers = ['指标', '本月实际', '上月实际', '环比变动', '去年同月', '同比变动', '预算值', '预算偏差']
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    # 行序故意打乱（人工在最前、中间插产量行；旧实现按 rows[4:7] 顺序会错绑）
    for label in ('其中：直接人工(元/盒)', '产量(盒)', '其中：直接材料(元/盒)', '其中：制造费用(元/盒)'):
        table.add_row().cells[0].text = label
    source = tmp_path / 'tpl.docx'
    doc.save(source)
    result = reports.install_template(source, 'monthly')
    installed = tmp_path / 'templates' / 'monthly.docx'
    doc2 = Document(installed)
    overview = next(t for t in doc2.tables if any('去年同月' in c.text for c in t.rows[0].cells))
    header_texts = [c.text.strip() for c in overview.rows[0].cells]
    yoy_col, rate_col = header_texts.index('去年同月'), header_texts.index('同比变动')
    def row_of(keyword):
        return next(r for r in overview.rows[1:] if keyword in r.cells[0].text)
    assert row_of('直接材料').cells[yoy_col].text == '{{去年材料成本}}'
    assert row_of('直接材料').cells[rate_col].text == '{{材料成本同比}}%'
    assert row_of('直接人工').cells[yoy_col].text == '{{去年人工成本}}'
    assert row_of('直接人工').cells[rate_col].text == '{{人工成本同比}}%'
    assert row_of('制造费用').cells[yoy_col].text == '{{去年制造费用}}'
    assert '同比补绑定 6 个单元格' in result['yoy_binding_note']
    # 占位符登记了补绑定的六个字段
    fields = {e['field'] for e in result['placeholders']}
    assert {'去年材料成本', '材料成本同比', '去年人工成本', '去年制造费用'} <= fields


def test_install_template_missing_yoy_columns_skips_binding(isolated_runtime, tmp_path, monkeypatch):
    import pharma.reports as reports
    from docx import Document
    monkeypatch.setattr(reports, 'RUNTIME_TEMPLATES', tmp_path / 'templates')
    doc = Document()
    for title in ('一、封面与基本信息', '二、总成本概览', '三、成本要素明细分析',
                  '四、重点产品专项分析', '五、对标分析', '六、总结与建议'):
        doc.add_paragraph(title)
    doc.add_paragraph('{{产品名称}}')
    headers = ['指标', '本月实际', '上月实际', '环比变动', '预算值']  # 无去年同月/同比列
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    table.add_row().cells[0].text = '其中：直接材料(元/盒)'
    source = tmp_path / 'tpl2.docx'
    doc.save(source)
    result = reports.install_template(source, 'monthly')
    assert '跳过' in result['yoy_binding_note']
    assert not any(e['field'].startswith('去年') for e in result['placeholders'])


# ---------- AUD-BENCH-01：allow_model=False 与只读缓存 ----------

def test_generate_allow_model_false_returns_rules_without_model_calls():
    from pharma.narrative import generate
    from pharma.metrics import analyze
    snapshot = analyze('中药一厂', '银黄口服液', '2026-05')
    result = generate(snapshot, [], allow_model=False)
    assert result['generation_mode'] == 'rules'
    assert result['model_responded'] is False
    assert result['findings']  # 确定性事实仍在（数字由程序绑定）
    # 无输入证据时结论显式为证据不足，而非模型解释
    assert all(f.get('origin') in (None, 'rules') for f in result['findings'])


def test_cached_generation_readonly_and_generate_write_visible(tmp_path):
    # 独立 gateway runtime：不读写真实 .runtime/model_gateway.sqlite3
    import pharma.narrative as narrative
    from pharma.metrics import analyze
    snapshot = analyze('中药一厂', '银黄口服液', '2026-05')
    evidence = {'status': 'PASS', 'knowledge_version': 'v-test', 'evidence': []}
    gateway = narrative.ModelGateway(runtime=tmp_path, key_file=str(tmp_path / 'k.key'))
    assert narrative.cached_generation(snapshot, evidence, gateway=gateway) is None  # 空缓存不命中且不抛错
    fake = {'status': 'PASS', 'model': 'stub', 'model_live': True, 'findings': [{'claim_type': 'hypothesis'}]}
    import sqlite3
    _sources, _excluded, _kv, gateway, key = narrative._prepare_generation(snapshot, evidence, gateway)
    with sqlite3.connect(gateway.dbpath) as db:
        db.execute('INSERT OR REPLACE INTO cache VALUES(?,?,?)',
                   (key, 12345.0, json.dumps(fake, ensure_ascii=False)))
    hit = narrative.cached_generation(snapshot, evidence, gateway=gateway)
    assert hit is not None and hit.get('cache_hit') is True and hit['model'] == 'stub'


def test_benchmark_async_default_queues_report_and_returns_rules_immediately(tmp_path, monkeypatch):
    """对标页默认 async（AUD-BENCH-01）：冷缓存即时返回规则解释 + 入队报告任务；
    入队失败不阻断响应；已完成任务的 narrative 被直接复用。"""
    from types import SimpleNamespace
    from pharma import api, narrative, context_services, knowledge
    from pharma.jobs import JobStore
    context = {'enterprise_id': 'competition', 'industry_id': 'pharmaceutical', 'data_snapshot': 'synthetic-only'}
    monkeypatch.setattr(api, 'selected_context', lambda _: 'pharmaceutical:competition')
    monkeypatch.setattr(api, 'benchmark_analysis',
                        lambda *a, **k: ({'snapshot_id': 's1', 'factory': 'A', 'product': 'P', 'month': '2026-05',
                                          'period': {'start': '2026-05', 'end': '2026-05'}, 'elements': []}, {'elements': []}))
    import pharma.industry as industry
    monkeypatch.setattr(industry, 'resolve_context',
                        lambda _: SimpleNamespace(model_dump=lambda: context, context_hash='scope'))
    monkeypatch.setattr(context_services, 'retrieve', lambda *a, **k: {'status': 'PASS', 'evidence': []})
    monkeypatch.setattr(knowledge, 'Knowledge', lambda *a, **k: SimpleNamespace(search=None))
    monkeypatch.setattr(api, 'store', JobStore(tmp_path / 'db'))
    enqueued = []
    def fake_enqueue(req,prepared_snapshot=None):
        assert prepared_snapshot is not None  # 后台解释必须绑定本次对标快照。
        job = {'id': 'job-async-1', 'status': 'QUEUED', 'result': {}}
        enqueued.append(req.model_dump())
        return job, {}
    monkeypatch.setattr(api, '_enqueue_report', fake_enqueue)
    def no_model(snapshot, evidence, use_cache=True, allow_model=True, gateway=None):
        assert allow_model is False, 'async 冷缓存路径不得发起模型调用'
        return {'status': 'DEGRADED', 'generation_mode': 'rules', 'findings': [{'claim_type': 'insufficient_evidence', 'rendered_text': '现有证据不足以确认差异原因。'}]}
    monkeypatch.setattr(narrative, 'generate', no_model)
    result = api.get_benchmark('P', '2026-05', 'A', 'B', context_id='pharmaceutical:competition')
    assert result['narrative']['model_status'] == 'QUEUED' and result['narrative']['job_id'] == 'job-async-1'
    assert result['narrative']['generation_mode'] == 'rules'
    assert enqueued and enqueued[0]['factory'] == 'A'
    # 显式只读基础分析不因冷缓存触发后台写入或付费请求。
    enqueued.clear()
    plain=api.get_benchmark('P','2026-05','A','B',context_id='pharmaceutical:competition',explain='none')
    assert plain['narrative']['generation_mode']=='rules' and not enqueued
    # 已完成任务 → 直接复用其 narrative，不再入队新任务
    enqueued.clear()
    monkeypatch.setattr(api, '_enqueue_report',
                        lambda req,prepared_snapshot=None: ({'id': 'job-done', 'status': 'SUCCEEDED',
                                      'result': {'narrative': {'status': 'PASS', 'generation_mode': 'llm', 'findings': []}}}, {}))
    result2 = api.get_benchmark('P', '2026-05', 'A', 'B', context_id='pharmaceutical:competition')
    assert result2['narrative']['generation_mode'] == 'llm' and 'model_status' not in result2['narrative']
    # 入队失败：响应保留规则解释并标注，不抛错
    def boom(req,prepared_snapshot=None):
        raise ValueError('NO_COMPLETE_PERIOD_DATA')
    monkeypatch.setattr(api, '_enqueue_report', boom)
    result3 = api.get_benchmark('P', '2026-05', 'A', 'B', context_id='pharmaceutical:competition')
    assert result3['narrative']['model_status'] == 'ENQUEUE_FAILED'
    assert result3['narrative']['generation_mode'] == 'rules'


# ---------- AUD-KB-01：knowledge.search 分块/适用性缓存与副本隔离 ----------

def test_knowledge_search_caches_chunks_and_applicability_with_isolation(tmp_path, monkeypatch):
    import pharma.knowledge as km
    monkeypatch.setattr(km, '_CHUNK_CACHE', {})
    monkeypatch.setattr(km, '_APPLICABILITY_CACHE', {})
    src = tmp_path / 'k.json'
    src.write_text(json.dumps([
        {'id': 'a', 'products': ['P1'], 'text': '银黄口服液配方含金银花与黄芩提取物，需核对用量与批次记录。'},
        {'id': 'b', 'products': ['P2'], 'text': '另一产品工艺说明：浓缩与制粒工序的温度控制要求完整记录。'},
    ], ensure_ascii=False))
    knowledge = km.Knowledge(root=tmp_path, source_files=(src,), vector_enabled=False)
    first = knowledge.search('配方 金银花', product='P1', mode='bm25')
    second = knowledge.search('配方 金银花', product='P1', mode='bm25')
    assert first['status'] == second['status'] == 'PASS'
    assert [e['evidence_id'] for e in first['evidence']] == [e['evidence_id'] for e in second['evidence']]
    # 命中路径：分块与适用性进入进程内缓存
    assert km._CHUNK_CACHE and km._APPLICABILITY_CACHE
    # 副本隔离：调用方修改返回的 evidence 内层结构不污染缓存
    if first['evidence'][0].get('products'):
        first['evidence'][0]['products'].append('POLLUTED')
    first['evidence'][0]['applicability']['reasons'].append('POLLUTED')
    third = knowledge.search('配方 金银花', product='P1', mode='bm25')
    assert all('POLLUTED' not in (e.get('products') or []) for e in third['evidence'])
    assert all('POLLUTED' not in e['applicability']['reasons'] for e in third['evidence'])
    # 不同过滤参数不串缓存：P2 只命中 b 分块
    other = knowledge.search('工艺 制粒', product='P2', mode='bm25')
    assert other['evidence'] and all('浓缩' in e['text'] for e in other['evidence'])
    # 索引重建（源变化→新版本指纹）后缓存自然失效出新结果
    src.write_text(json.dumps([
        {'id': 'a', 'products': ['P1'], 'text': '银黄口服液配方含金银花与黄芩提取物，需核对用量与批次记录。'},
        {'id': 'b', 'products': ['P2'], 'text': '另一产品工艺说明：浓缩与制粒工序的温度控制要求完整记录。'},
        {'id': 'c', 'products': ['P1'], 'text': '新增分块：金银花市场行情上涨说明与采购核查要求逐条记录。'},
    ], ensure_ascii=False))
    knowledge2 = km.Knowledge(root=tmp_path, source_files=(src,), vector_enabled=False)
    knowledge2.build()  # 真实链路（worker）检索前总是幂等 build；search 仅在术语变化时自动重建
    refreshed = knowledge2.search('行情 上涨', product='P1', mode='bm25')
    assert any('行情' in e['text'] for e in refreshed['evidence'])


# ---------- AUD-IMP-01/02/06/03/07（第二审查报告 FIX-A，2026-09-24） ----------

def _contest_summary_csv():
    """题包《成本汇总》同构宽表（单位成本列 元/盒 口径，数据自洽）。"""
    return ('工厂,产品名称,产品规格,月份,产量(盒),直接材料(元/盒),直接人工(元/盒),制造费用(元/盒),单位成本(元/盒),总成本(元)\n'
            '中药一厂,银黄口服液,10ml×10支/盒,2026-05,45000,6.50,2.10,2.10,10.70,481500\n'
            '中药一厂,银黄口服液,10ml×10支/盒,2026-04,44000,6.20,2.05,2.05,10.30,453200\n')


def test_contest_summary_import_converts_unit_cost_and_totals_match(isolated_runtime, tmp_path):
    """AUD-IMP-01 验收：题包成本汇总经 UI 同款链路导入——单位成本列按产量换算，
    发布后 analyze 的总成本=481500、单位成本=10.70，与真值一致。"""
    record = data_import.create_upload('business', '成本汇总.csv', _contest_summary_csv().encode('utf-8'), 'cost_summary')
    mapping = _mapping_for(record)
    result = data_import.validate_business(record, mapping, {})
    assert result['status'] == 'VALID', result['errors'][:3]
    # 不变量（AUD-IMP-02）：Σ要素（换算后）与总成本列勾稽通过
    published = data_import.publish_business(record, mapping, {}, '审计量纲企业', 'pharmaceutical', '盒')
    assert published['dataset_facts'] > 0
    import pharma.industry as industry
    context = industry.resolve_context(published['context_id'])
    snapshot = industry.analyze_reference(published['context_id'], factory='中药一厂',
                                          product='银黄口服液', month='2026-05')
    assert Decimal(snapshot['metrics']['unit_cost']['value']) == Decimal('10.70')
    assert Decimal(snapshot['metrics']['total_cost']['value']) == Decimal('481500')


def test_unit_cost_without_quantity_is_invalid(isolated_runtime):
    record = data_import.create_upload('business', '无产量汇总.csv',
        ('工厂,产品名称,月份,直接材料(元/盒),直接人工(元/盒),制造费用(元/盒)\n'
         '甲厂,P1,2026-01,6.50,2.10,2.10\n').encode('utf-8'), 'cost_summary')
    result = data_import.validate_business(record, _mapping_for(record), {})
    assert result['status'] == 'INVALID'
    assert any('单位成本列' in e['reason'] and '产量' in e['reason'] for e in result['errors'])


def test_total_cost_mismatch_rejected(isolated_runtime):
    """AUD-IMP-02 验收：总成本列与 Σ要素 不一致（勾稽破坏）→ INVALID 带定位。"""
    bad = ('工厂,产品名称,月份,产量(盒),直接材料(元/盒),直接人工(元/盒),制造费用(元/盒),总成本(元)\n'
           '甲厂,P1,2026-01,100,3.00,1.00,1.00,999\n')  # Σ=500 vs 999
    record = data_import.create_upload('business', '勾稽不符.csv', bad.encode('utf-8'), 'cost_summary')
    result = data_import.validate_business(record, _mapping_for(record), {})
    assert result['status'] == 'INVALID'
    assert any('总成本列' in e['reason'] and '要素合计' in e['reason'] for e in result['errors'])


def test_duplicate_summary_rows_hard_fail_but_detail_accumulates(isolated_runtime):
    """AUD-IMP-06：汇总表重复行 INVALID；明细长表同要素多行累加仍合法。"""
    dup = ('工厂,产品名称,月份,产量(盒),直接材料(元/盒),直接人工(元/盒),制造费用(元/盒),总成本(元)\n'
           '甲厂,P1,2026-01,100,3.00,1.00,1.00,500\n'
           '甲厂,P1,2026-01,100,3.00,1.00,1.00,500\n')
    record = data_import.create_upload('business', '重复汇总.csv', dup.encode('utf-8'), 'cost_summary')
    result = data_import.validate_business(record, _mapping_for(record), {})
    assert result['status'] == 'INVALID'
    assert any('重复出现' in e['reason'] for e in result['errors'])
    detail = ('工厂,产品名称,月份,产量(盒),原材料名称,原材料总成本(元)\n'
              '甲厂,P1,2026-01,100,金银花,200\n'
              '甲厂,P1,2026-01,100,黄芩,100\n')
    record2 = data_import.create_upload('business', '材料明细.csv', detail.encode('utf-8'), 'material_detail')
    detail_mapping = {'工厂': 'factory_id', '产品名称': 'product_id', '月份': 'period',
                      '产量(盒)': 'quantity', '原材料总成本(元)': 'element:material'}
    result2 = data_import.validate_business(record2, detail_mapping, {'scenario': 'actual'})
    assert result2['status'] == 'VALID', result2['errors'][:2]
    assert result2['statistics']['cost_cells'] == 1  # 同要素键累加为一个金额键（合法形态）


def test_reimport_different_data_keeps_old_context(isolated_runtime, tmp_path):
    """AUD-IMP-03 验收：同名企业重导不同数据 → 新目录新上下文，旧注册条目数据不被覆盖。"""
    import pharma.industry as industry
    csv_a = _contest_summary_csv()
    record_a = data_import.create_upload('business', 'a.csv', csv_a.encode('utf-8'), 'cost_summary')
    first = data_import.publish_business(record_a, _mapping_for(record_a), {}, '同名家', 'pharmaceutical', '盒')
    csv_b = csv_a.replace('481500', '482000').replace('10.70,481500', '10.70,482000').replace('453200', '454000')
    record_b = data_import.create_upload('business', 'b.csv', csv_b.encode('utf-8'), 'cost_summary')
    second = data_import.publish_business(record_b, _mapping_for(record_b), {}, '同名家', 'pharmaceutical', '盒')
    assert first['enterprise_id'] != second['enterprise_id']  # 数据指纹不同 → 不同目录
    snap_old = industry.analyze_reference(first['context_id'], factory='中药一厂', product='银黄口服液', month='2026-05')
    assert Decimal(snap_old['metrics']['total_cost']['value']) == Decimal('481500')  # 旧数据未被覆盖


def test_publish_knowledge_docx_includes_tables(isolated_runtime, tmp_path):
    """AUD-IMP-07 验收：含表格的 docx 知识导入后可检到表格内文本。"""
    from docx import Document
    doc = Document()
    doc.add_paragraph('设备维护规程正文段落，说明巡检要求。')
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = '项目'; table.rows[0].cells[1].text = '要求'
    table.rows[1].cells[0].text = '计量盘'; table.rows[1].cells[1].text = '每季度更换并记录台账'
    src = tmp_path / '设备表.docx'
    doc.save(src)
    record = data_import.create_upload('knowledge', '设备表.docx', src.read_bytes(), 'enterprise')
    result = data_import.publish_knowledge(record, {'title': '设备'})
    assert '计量盘' in (tmp_path / 'imports' / record['id'] / 'knowledge_entry.json').read_text(encoding='utf-8') or result['characters'] > 30
    entry = json.loads((data_import.IMPORTS_ROOT / record['id'] / 'knowledge_entry.json').read_text(encoding='utf-8'))
    assert '每季度更换' in entry['text']  # 表格内容进入知识文本


def test_attribution_failure_degrades_not_fails_batch(isolated_runtime, tmp_path, monkeypatch):
    """AUD-IMP-04 验收：发布成功后归因概览抛错 → 批次仍 SUCCEEDED，提示降级。"""
    import pharma.import_pipeline as pipeline
    from pharma.jobs import JobStore
    store = JobStore(tmp_path / 'db.sqlite3')
    record = data_import.create_upload('business', '汇总.csv', _contest_summary_csv().encode('utf-8'), 'cost_summary')
    job = store.enqueue('data_parse', {'import_ids': [record['id']], 'enterprise_name': '归因降级企业', 'quantity_unit': '盒'})
    def boom(context_id):
        raise RuntimeError('归因模拟故障')
    monkeypatch.setattr(pipeline, '_attribution_overview', boom)
    pipeline.run_data_parse(store, store.get(job['id']))
    final = store.get(job['id'])
    assert final['status'] == 'SUCCEEDED', final.get('error')
    assert final['result']['attribution']['attribution_mode'] == 'unavailable'
    assert data_import.get_import(record['id'])['status'] == 'PARSED'


def test_parsing_records_recoverable_after_stale(isolated_runtime, monkeypatch):
    """AUD-IMP-05 验收：PARSING 超时记录可被 waiting_imports 重新领取。"""
    from datetime import datetime, timedelta, timezone
    record = data_import.create_upload('business', '卡死.csv', _contest_summary_csv().encode('utf-8'), 'cost_summary')
    data_import.mark_import_status(record, 'PARSING')
    assert data_import.waiting_imports('business') == []  # 刚标记（新鲜）不可重领
    # 把 updated 改成 31 分钟前
    with data_import._connect() as db:
        db.execute('UPDATE imports SET updated=? WHERE id=?',
                   ((datetime.now(timezone.utc) - timedelta(seconds=1900)).isoformat(), record['id']))
    recovered = data_import.waiting_imports('business')
    assert [r['id'] for r in recovered] == [record['id']]


# ---------- AUD-RAG-01/REP-01（第二审查报告，2026-09-24） ----------

def test_competition_context_knowledge_includes_supplement(tmp_path, monkeypatch):
    """RAG-01 验收：竞赛上下文（报告/对标链构造方式）也纳入补充知识，
    此前只有交互式 Knowledge() 才挂 extra_dir。"""
    import pharma.knowledge as km
    pkg_dir = tmp_path / '00_赛题原始资料/模拟数据_V1.1_净化解压/创灵境_考题模拟数据/03_制药知识文档'
    pkg_dir.mkdir(parents=True)
    (pkg_dir / '配方A.txt').write_text('产品配方：金银花用量与提取工艺要求，投料须与批记录一致。', encoding='utf-8')
    sup = tmp_path / 'supplement'; sup.mkdir()
    (sup / '异常处理记录.txt').write_text('历史成本异常处理记录：金银花采购价上涨时核查调价条款与库存结构。', encoding='utf-8')
    monkeypatch.setattr('pharma.config.KNOWLEDGE_SUPPLEMENT_DIR', sup)
    context = {'industry_id': 'pharmaceutical', 'enterprise_id': 'competition'}
    k = km.Knowledge(root=tmp_path, vector_enabled=False, context=context)
    assert k.extra_dir == sup  # 修复前：竞赛上下文 extra_dir 为 None
    manifest = k.build()
    names = set(manifest['sources'])
    assert any('异常处理' in str(n) for n in names) and any('配方A' in str(n) for n in names)
    hit = k.search('异常 调价条款', mode='bm25', limit=3)
    assert any('异常处理' in str(e.get('source')) for e in hit['evidence'])
    # 行业包 allowlist 语义不变（显式 source_files 不追加）
    via_files = km.Knowledge(root=tmp_path, vector_enabled=False,
                             context={'industry_id': 'mechanical'}, source_files=(pkg_dir / '配方A.txt',))
    assert via_files.extra_dir is None


def test_explanation_presence_accepts_legacy_summary_section(tmp_path):
    """REP-01 验收：legacy findings（section=summary）的模型解释在正文
    任意位置存在时不再判缺失（此前直接 FAILED）。"""
    from docx import Document
    from pharma.reports import explanation_presence
    doc = Document()
    doc.add_paragraph('一、封面与基本信息')
    doc.add_paragraph('单位成本环比上升的主要解释文本出现在这里。')
    path = tmp_path / 'r.docx'
    doc.save(path)
    narrative = {'findings': [{'origin': 'model', 'claim_type': 'hypothesis',
                               'section': 'summary',
                               'rendered_text': '单位成本环比上升的主要解释文本出现在这里。'}]}
    result = explanation_presence(path, narrative)
    assert result['explanation_binding_failures'] == []
    missing = {'findings': [{'origin': 'model', 'claim_type': 'hypothesis', 'section': 'summary',
                             'rendered_text': '这段话不在正文里，应仍判缺失。'}]}
    assert explanation_presence(path, missing)['explanation_binding_failures']


# ---------- AUD-TST-03：题包真数据独立金标（期望值硬编码，不经引擎推导） ----------

def test_golden_metrics_yinhuang_2026_05():
    """金标值由审计独立复算脚本从题包原始 CSV 以 Decimal 手工推导
    （docs/audits/20260923-1f78b2f/EVIDENCE/verify_metrics.py，2026-09-23）。
    引擎口径若被改坏（贡献度公式/比较基期/预算桥），本测试变红——
    弥补 verify_docx 用 build_bindings 自产期望的自洽盲区。"""
    from pharma.metrics import analyze, benchmark
    snap = analyze('中药一厂', '银黄口服液', '2026-05')
    assert Decimal(snap['metrics']['unit_cost']['value']) == Decimal('11.21')
    assert abs(Decimal(snap['comparison']['mom']['rate']) - Decimal('2.844036697247706422018348623853211009174')) < Decimal('1e-20')
    assert abs(Decimal(snap['comparison']['yoy']['rate']) - Decimal('4.668534080298786181139122315592903828198')) < Decimal('1e-20')
    assert abs(Decimal(snap['comparison']['budget']['rate']) - Decimal('5.754716981132075471698113207547169811321')) < Decimal('1e-20')
    contributions = {e['key']: Decimal(e['comparisons']['mom']['unit']['contribution']) for e in snap['elements']}
    assert abs(contributions['materials'] - Decimal('70.96774193548387096774193548387096774194')) < Decimal('1e-20')
    assert abs(contributions['labor'] - Decimal('9.677419354838709677419354838709677419355')) < Decimal('1e-20')
    assert abs(contributions['overhead'] - Decimal('19.35483870967741935483870967741935483871')) < Decimal('1e-20')
    assert snap['alerts'] == []  # 本月三要素环比均未严格超过±10%（金标判定）
    bridge = snap['budget_bridge']
    assert Decimal(bridge['quantity_effect']) == Decimal('84800.0')
    assert Decimal(bridge['unit_cost_effect']) == Decimal('35380.00')
    assert Decimal(bridge['total_delta']) == Decimal('120180.0')
    comp = benchmark('银黄口服液', '2026-05', left='中药一厂', right='中药二厂')
    assert Decimal(comp['summary'][0]['delta']) == Decimal('-0.39')
    assert abs(Decimal(comp['summary'][0]['rate']) - Decimal('-3.362068965517241379310344827586206896552')) < Decimal('1e-20')
