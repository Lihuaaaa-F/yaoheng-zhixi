"""Independent tampering tests for visible report metric-role bindings."""
from docx import Document
import pytest

from pharma.industry import analyze_reference, benchmark_reference
from pharma.reference_report import render, verify
from pharma.reports import layout_text


def make_report(tmp_path, context='mechanical_demo:synthetic-mechanical', benchmark=False):
    if benchmark:
        snapshot, comparison = benchmark_reference(context, 'DEMO-01', '2026-06', '示范工厂A', '示范工厂B')
    else:
        snapshot, comparison = analyze_reference(context, month='2026-06'), None
    # Chart rendering is orthogonal to the visible table/paragraph contract.
    snapshot['trend'] = []
    output = tmp_path / 'report.docx'
    result = render(snapshot, {'findings': []}, {'evidence': []}, output, comparison)
    return snapshot, output, result


@pytest.mark.parametrize('context', [
    'mechanical_demo:synthetic-mechanical',
    'chemical_demo:synthetic-chemical',
    'pharmaceutical:synthetic-pharma',
])
def test_legitimate_render_is_role_bound(tmp_path, context):
    snapshot, output, rendered = make_report(tmp_path, context)
    result = verify(output, snapshot)
    assert result['status'] == rendered['verification']['status'] == 'PASS'
    assert not result['numeric_binding_failures']
    # Only the exact cells and paragraphs actually inspected are counted.
    expected = 6 + 3 * len(snapshot['elements']) + len(snapshot['alerts']) + len(snapshot['specialized_metrics'])
    assert result['numeric_bindings_checked'] == expected


def test_core_numbers_cannot_be_swapped_or_repaired_in_appendix(tmp_path):
    snapshot, output, _ = make_report(tmp_path)
    doc = Document(output)
    table = doc.tables[0]
    assert [layout_text(row.cells[1].text) for row in table.rows[1:]] == ['3600.00 元', '120.00 件', '30.00 元/件']
    table.rows[1].cells[1].text = '30.00 元'
    table.rows[2].cells[1].text = '3600.00 件'
    table.rows[3].cells[1].text = '120.00 元/件'
    doc.add_paragraph('附录正确数字：3600.00 元；120.00 件；30.00 元/件')
    doc.save(output)
    result = verify(output, snapshot)
    assert result['status'] == 'FAIL'
    assert {f['binding'] for f in result['numeric_binding_failures']} == {
        'core:total_cost', 'core:quantity', 'core:unit_cost'}


@pytest.mark.parametrize('replacement', ['单位产品机时：99999.00 小时/件', '单位产品机时：0.50 分钟/件'])
def test_specialized_metric_value_and_unit_are_checked(tmp_path, replacement):
    snapshot, output, _ = make_report(tmp_path)
    doc = Document(output)
    target = next(p for p in doc.paragraphs if p.text.startswith('单位产品机时：'))
    assert layout_text(target.text) == '单位产品机时：0.50 小时/件'
    target.text = replacement
    doc.add_paragraph('单位产品机时：0.50 小时/件')
    doc.save(output)
    result = verify(output, snapshot)
    assert result['status'] == 'FAIL'
    assert [f['binding'] for f in result['numeric_binding_failures']] == ['specialized:machine_hours_per_piece']


def test_same_numbers_at_wrong_location_or_role_fail(tmp_path):
    snapshot, output, _ = make_report(tmp_path)
    doc = Document(output)
    # Move the complete, numerically correct overview to the final section.
    table = doc.tables[0]
    doc.element.body.insert(len(doc.element.body) - 1, table._tbl)
    doc.save(output)
    result = verify(output, snapshot)
    assert result['status'] == 'FAIL'
    assert len(result['numeric_binding_failures']) == 3


@pytest.mark.parametrize('target', ['comparison', 'element', 'core_unit', 'duplicate'])
def test_other_visible_numeric_bindings_fail_closed(tmp_path, target):
    snapshot, output, _ = make_report(tmp_path)
    doc = Document(output)
    if target == 'comparison':
        next(p for p in doc.paragraphs if p.text.startswith('环比：')).text = '环比：99999.00 %'
    elif target == 'element':
        doc.tables[1].rows[1].cells[1].text = '99999.00'
    elif target == 'core_unit':
        doc.tables[0].rows[3].cells[1].text = '30.00 元/kg'
    else:
        from copy import deepcopy
        row = doc.tables[0].rows[1]._tr
        row.addnext(deepcopy(row))
    doc.save(output)
    result = verify(output, snapshot)
    assert result['status'] == 'FAIL'
    assert len(result['numeric_binding_failures']) == 1


def test_benchmark_table_is_checked_against_bound_comparison(tmp_path):
    snapshot, output, rendered = make_report(tmp_path, benchmark=True)
    assert rendered['verification']['status'] == 'PASS'
    doc = Document(output)
    doc.tables[2].rows[1].cells[3].text = '99999.00'
    doc.save(output)
    result = verify(output, snapshot)
    assert result['status'] == 'FAIL'
    assert len(result['numeric_binding_failures']) == 1
    assert result['numeric_binding_failures'][0]['binding'].startswith('benchmark:')


def test_undefined_comparison_has_no_percent_and_month_is_not_repeated(tmp_path):
    _, output, _ = make_report(tmp_path)
    lines=[layout_text(p.text) for p in Document(output).paragraphs]
    assert any(line.startswith('同比：N/A') and '%' not in line for line in lines)
    assert all('2026-06 至 2026-06' not in line for line in lines)


def test_report_uses_frozen_template_without_loading_live_template(tmp_path, monkeypatch):
    import json
    from pathlib import Path
    from pharma.industry import load_pack, PACKS
    snapshot=analyze_reference('mechanical_demo:synthetic-mechanical',month='2026-06')
    pack=load_pack('mechanical_demo')
    snapshot['report_template']=json.loads((PACKS/pack.id/pack.template_entry).read_text())
    snapshot['trend']=[]
    original=Path.read_text
    def read(self,*args,**kwargs):
        if self.name=='template.json':raise AssertionError('live template must not rebind a saved snapshot')
        return original(self,*args,**kwargs)
    monkeypatch.setattr(Path,'read_text',read)
    output=tmp_path/'frozen.docx'
    assert render(snapshot,{'findings':[]},{'evidence':[]},output)['status']=='PASS'
    assert verify(output,snapshot)['status']=='PASS'


def test_task_blocks_and_source_heading_keep_with_next(tmp_path):
    snapshot=analyze_reference('mechanical_demo:synthetic-mechanical',month='2026-06');snapshot['trend']=[]
    finding={'rendered_text':'缺少经签署的机时记录，不能归因。','suggestion':'核对机时记录',
             'verification_target':'机时记录','responsible_role':'生产主管','deadline_basis':'月结前',
             'expected_evidence':['经签署的机时记录']}
    output=tmp_path/'tasks.docx'
    render(snapshot,{'findings':[finding]},{'evidence':[{'source':'合成来源','location':'记录1'}]},output)
    doc=Document(output)
    for prefix in ('核查对象：','核查行动：','责任岗位：','期限依据：','证据来源'):
        assert next(p for p in doc.paragraphs if p.text.startswith(prefix)).paragraph_format.keep_with_next
    assert next(p for p in doc.paragraphs if p.text.startswith('预期证据：')).paragraph_format.keep_together
