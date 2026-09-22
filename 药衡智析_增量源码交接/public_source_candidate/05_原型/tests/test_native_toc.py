"""原生目录与版式修复回归（2026-09-22 用户四项反馈）。

钉住的合同：
① 目录是 Word 原生 TOC 域（begin/separate/end + dirty），缓存条目用
   w:hyperlink 指向标题书签——未更新域也可点击跳转，Word 打开时自动重建；
② 小标题(x.y)挂 Heading 2 进目录，x.y.z 挂 Heading 3 不进目录；字号层级
   H1 小三15 / H2 四号14 / H3 小四12；
③ 题包原件自带的阅读指南截断（"成本要素明、重点产品专、"）渲染期补全；
④ 数字-单位防断行只用 NBSP，不再注入 U+FEFF（Word 渲染为可见异常符号）。
"""
import pytest
from docx import Document
from docx.oxml.ns import qn
from pharma import reports


def _mini_doc():
    doc = Document()
    doc.add_paragraph('目     录')
    doc.add_paragraph('一、封面与基本信息')
    doc.add_paragraph('正文段落，本月 7.26 元/盒。')
    doc.add_paragraph('2.1 核心指标一览')
    doc.add_paragraph('3.1.1 原材料成本明细')
    return doc


def _heads(doc):
    heads = [(p, reports.layout_text(p.text).strip()) for p in doc.paragraphs]
    return [(p, t) for p, t in heads if reports._heading_level(t)]


def test_native_toc_field_structure_and_dirty_flag():
    doc = _mini_doc()
    title = doc.paragraphs[0]
    heads = _heads(doc)
    reports._ensure_outline_styles(doc)
    reports._apply_outline_styles(doc, heads)
    reports._insert_native_toc(doc, title, heads)
    body = doc.element.body
    flds = body.findall('.//' + qn('w:fldChar'))
    assert [f.get(qn('w:fldCharType')) for f in flds] == ['begin', 'separate', 'end']
    assert flds[0].get(qn('w:dirty')) == 'true'  # Word 打开时自动更新
    instr = body.findall('.//' + qn('w:instrText'))[0].text
    assert 'TOC' in instr and '\\h' in instr and '\\o "1-2"' in instr


def test_toc_entries_hyperlinked_and_level_filtered():
    doc = _mini_doc()
    title = doc.paragraphs[0]
    heads = _heads(doc)
    reports._ensure_outline_styles(doc)
    reports._apply_outline_styles(doc, heads)
    reports._insert_native_toc(doc, title, heads)
    hyperlinks = doc.element.body.findall('.//' + qn('w:hyperlink'))
    anchors = [h.get(qn('w:anchor')) for h in hyperlinks]
    # 一级章 + 2.1 小标题两条入目录；3.1.1 为三级不进
    assert anchors == ['YH_SEC_T0', 'YH_SEC_T1']
    texts = [''.join(t.text or '' for t in h.findall('.//' + qn('w:t'))) for h in hyperlinks]
    assert any(t.lstrip('\u3000').startswith('2.1') for t in texts)      # 小标题进目录
    assert not any('3.1.1' in t for t in texts)                          # 三级不进目录
    assert all(t.rstrip().endswith('第1页') for t in texts)               # 页码占位待 convert_pdf 回填


def test_heading_styles_and_size_hierarchy():
    doc = _mini_doc()
    heads = _heads(doc)
    reports._ensure_outline_styles(doc)
    reports._apply_outline_styles(doc, heads)
    by_text = {p.text.strip(): p for p in doc.paragraphs}
    assert by_text['一、封面与基本信息'].style.name == 'Heading 1'
    assert by_text['2.1 核心指标一览'].style.name == 'Heading 2'
    assert by_text['3.1.1 原材料成本明细'].style.name == 'Heading 3'
    sizes = {p.runs[0].font.size.pt for p in by_text.values() if p.runs and p.text.strip() in
             ('一、封面与基本信息', '2.1 核心指标一览', '3.1.1 原材料成本明细')}
    assert sizes == {15.0, 14.0, 12.0}  # 小三/四号/小四


def test_reading_guide_truncation_is_repaired():
    doc = Document()
    cell = doc.add_table(rows=1, cols=1).rows[0].cells[0]
    cell.text = '三、成本要素明、四、重点产品专、五、对标分析、六、总结与建议'
    reports._fix_reading_guide(doc)
    assert cell.text == '三、成本要素明细分析、四、重点产品专项分析、五、对标分析、六、总结与建议'


def test_number_units_use_nbsp_never_ufeff():
    doc = _mini_doc()
    para = doc.paragraphs[2]
    reports.protect_number_units(para)
    text = ''.join(t.text or '' for t in para._p.iter('{' + reports.W + '}t'))
    assert '\ufeff' not in text
    assert '7.26\u00a0元' in text  # NBSP 绑定数字与单位（Word 显示为普通空格）
