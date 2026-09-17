"""Pack-configured report composition for independent synthetic datasets.
Shares the report compiler's typography, conversion and artifact contracts.
"""
from pathlib import Path
import hashlib,json,re
from docx import Document
from .reports import style_reader,number,RESIDUAL,layout_text,benchmark_labels

HEADINGS=['一、封面与基本信息','二、总成本概览','三、成本要素明细分析','四、重点产品专项分析','五、对标分析','六、总结与建议']

def verify(path, snapshot, benchmark=None):
    """Check visible values at their semantic role and section, not anywhere in XML.

    Each binding is one exact visible cell or paragraph comparison, including its
    display precision and unit. Missing/duplicate roles fail; appendix numbers
    cannot satisfy a binding in a different section.
    """
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from .industry import load_pack, PACKS
    pack = load_pack(snapshot['analysis_context']['industry_id'])
    template = snapshot.get('report_template') or json.loads((PACKS / pack.id / pack.template_entry).read_text())
    headings = template['sections']
    doc = Document(path)
    sections = {i: {'paragraphs': [], 'tables': []} for i in range(len(headings))}
    observed, section, all_text = [], None, []
    for node in doc.element.body:
        if node.tag == qn('w:p'):
            paragraph = Paragraph(node, doc)
            text = layout_text(paragraph.text)
            all_text.append(text)
            if paragraph.style.name == 'Heading 1':
                observed.append(text)
                section = headings.index(text) if text in headings else None
            elif section is not None:
                sections[section]['paragraphs'].append(text)
        elif node.tag == qn('w:tbl'):
            rows = [[layout_text(cell.text) for cell in row.cells] for row in Table(node, doc).rows]
            all_text.extend(cell for row in rows for cell in row)
            if section is not None:
                sections[section]['tables'].append(rows)
    failures, checked = [], []

    def bind(key, actual, expected):
        checked.append(key)
        if actual != [expected]:
            failures.append({'binding': key, 'expected': expected, 'actual': actual,
                             'reason': 'missing_or_duplicate_role' if len(actual) != 1 else 'value_unit_or_format_mismatch'})

    def table_value(index, headers, label, column):
        tables = [rows for rows in sections.get(index, {}).get('tables', []) if rows and rows[0] == headers]
        if len(tables) != 1:
            return []
        return [row[column] for row in tables[0][1:] if len(row) == len(headers) and row[0] == label]

    def paragraph_value(index, prefix):
        return [text for text in sections.get(index, {}).get('paragraphs', []) if text.startswith(prefix)]

    m = snapshot['metrics']
    for key, label in [('total_cost', '总成本'), ('quantity', '合格产出'), ('unit_cost', '单位成本')]:
        bind('core:' + key, table_value(1, ['指标', '确定性结果'], label, 1),
             number(m[key]) + ' ' + m[key]['unit'])
    for key, label in [('mom', '环比'), ('yoy', '同比'), ('budget', '预算差异')]:
        value = m[key]
        expected = label + '：' + number(value) + (' ' + value['unit'] if value.get('value') is not None else '')
        if value.get('reason'):
            expected += '；' + value['reason']
        bind('comparison:' + key, paragraph_value(1, label + '：'), expected)
    for element in snapshot['elements']:
        for column, key in enumerate(('total', 'unit', 'unit_mom'), 1):
            bind('element:' + element['key'] + ':' + key,
                 table_value(2, ['要素', '总成本', '单位成本', '单位环比（%）'], element['name'], column),
                 number(element.get(key)))
    for alert in snapshot['alerts']:
        prefix = alert['element'] + ' · ' + ('单位成本' if alert['basis'] == 'unit' else '总成本') + '环比 '
        bind('alert:' + alert['alert_id'], paragraph_value(2, prefix),
             prefix + number(alert['rate']) + '%，超过合成演示阈值。')
    for key in pack.strategies:
        label = {'machine_hours_per_piece': '单位产品机时', 'energy_per_kg': '单位合格产出能耗'}.get(key, key)
        bind('specialized:' + key, paragraph_value(3, label + '：'),
             label + '：' + number(m[key]) + ' ' + m[key]['unit'])
    comparison = benchmark if benchmark is not None else snapshot.get('benchmark_context', {})
    left, right, direction = benchmark_labels(comparison)
    benchmark_unit = snapshot['metrics']['unit_cost']['unit']
    benchmark_headers = ['要素', left+'（'+benchmark_unit+'）', right+'（'+benchmark_unit+'）', '差异（'+benchmark_unit+'）']
    for element in (comparison or {}).get('elements', []):
        for column, key in enumerate(('left', 'right', 'delta'), 1):
            bind('benchmark:' + element['key'] + ':' + key,
                 table_value(4, benchmark_headers, element.get('name', element['key']), column),
                 number(element.get(key)))
    residual = RESIDUAL.findall('\n'.join(all_text))
    headings_valid = len(headings) == 6 and observed == headings
    core = not any(item['binding'].startswith('core:') for item in failures)
    return {'status': 'PASS' if not failures and not residual and headings_valid else 'FAIL',
            'core_numbers': core, 'residual_placeholders': residual,
            'headings': observed, 'headings_valid': headings_valid,
            'numeric_bindings_checked': len(checked), 'numeric_binding_failures': failures,
            'contract': 'generic-v2-role-bound', 'scope': '表格及段落的角色、位置、数值和显示单位绑定；图像数值及真人评分不在此检查内'}

def render(snapshot,narrative,evidence,output,benchmark=None):
    from .narrative import render_visible_text
    from .industry import load_pack,PACKS
    pack=load_pack(snapshot['analysis_context']['industry_id'])
    template=snapshot.get('report_template') or json.loads((PACKS/pack.id/pack.template_entry).read_text());headings=template['sections']
    if len(headings)!=6:raise ValueError('REPORT_SECTION_CONTRACT_REQUIRES_SIX')
    doc=Document();m=snapshot['metrics'];unit=m['unit_cost']['unit'];quantity_unit=snapshot['quantity_unit']
    def paragraph(text):
        if RESIDUAL.search(text):raise ValueError('UNRESOLVED_REPORT_FIELD')
        return doc.add_paragraph(text)
    def table(headers,rows):
        t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
        for c,x in zip(t.rows[0].cells,headers):c.text=str(x)
        for row in rows:
            for c,x in zip(t.add_row().cells,row):c.text=str(x)
    doc.add_heading(pack.name+' · '+snapshot['product']+'成本分析报告',0)
    paragraph(template['required_notice'])
    for index,heading in enumerate(headings):
        doc.add_heading(heading,1)
        if index==0:
            period_label=snapshot['period']['start'] if snapshot['period']['start']==snapshot['period']['end'] else snapshot['period']['start']+' 至 '+snapshot['period']['end']
            paragraph(snapshot['factory']+' · '+period_label)
            paragraph('规格：'+snapshot['specification']+'；币种：'+snapshot['currency']+'；口径：完工产出。')
        elif index==1:
            table(['指标','确定性结果'],[['总成本',number(m['total_cost'])+' '+m['total_cost']['unit']],['合格产出',number(m['quantity'])+' '+quantity_unit],['单位成本',number(m['unit_cost'])+' '+unit]])
            for label in ('mom','yoy','budget'):
                val=m[label];paragraph({'mom':'环比','yoy':'同比','budget':'预算差异'}[label]+'：'+number(val)+(' '+val['unit'] if val.get('value') is not None else '')+('；'+val['reason'] if val.get('reason') else ''))
            paragraph('季度采用期间成本总和除以独立产量总和；缺失基期或零分母保持无定义。')
        elif index==2:
            table(['要素','总成本','单位成本','单位环比（%）'],[[e['name'],number(e['total']),number(e['unit']),number(e.get('unit_mom'))] for e in snapshot['elements']])
            for alert in snapshot['alerts']:paragraph(alert['element']+' · '+('单位成本' if alert['basis']=='unit' else '总成本')+'环比 '+number(alert['rate'])+'%，超过合成演示阈值。')
            if snapshot['trend']:
                import matplotlib
                matplotlib.use('Agg')
                import matplotlib.pyplot as plt
                from matplotlib.font_manager import FontProperties
                from .config import APP
                font=FontProperties(fname=str(APP/'assets/fonts/NotoSansSC-Regular.ttf'))
                fig,ax=plt.subplots(figsize=(7,2.3));ax.plot([r['month'] for r in snapshot['trend']],[float(r['unit_cost']) if r['unit_cost'] is not None else float('nan') for r in snapshot['trend']],marker='o',color='#246479');ax.set_ylabel(unit,fontproperties=font);ax.set_title('合成演示 · 单位成本趋势',fontproperties=font);fig.tight_layout()
                image=Path(output).with_suffix('.png');image.parent.mkdir(parents=True,exist_ok=True);fig.savefig(image,dpi=150);plt.close(fig)
                from docx.shared import Cm
                doc.add_picture(str(image),width=Cm(17))
        elif index==3:
            for key in pack.strategies:paragraph({'machine_hours_per_piece':'单位产品机时','energy_per_kg':'单位合格产出能耗'}.get(key,key)+'：'+number(m[key])+' '+m[key]['unit'])
            paragraph('专用指标以明确驱动事实计算，数值和阈值均为合成假定，不是行业基准。')
            for finding in narrative.get('findings',[]):
                paragraph(render_visible_text(finding.get('rendered_text') or finding.get('text_template',''),snapshot,evidence.get('evidence',[]),finding.get('metric_refs',[])))
        elif index==4:
            if benchmark and benchmark.get('elements'):
                left,right,direction=benchmark_labels(benchmark)
                paragraph(direction)
                table(['要素',left+'（'+unit+'）',right+'（'+unit+'）','差异（'+unit+'）'],[[e.get('name',e.get('key','')),number(e.get('left')),number(e.get('right')),number(e.get('delta'))] for e in benchmark['elements']])
            else:paragraph('未选择跨厂比较；不构造明细或归因。')
            paragraph(snapshot['details']['reason'])
        else:
            for f in narrative.get('findings',[]):
                if not f.get('suggestion'):continue
                for key,label in [('verification_target','核查对象'),('suggestion','核查行动'),('responsible_role','责任岗位'),('deadline_basis','期限依据')]:
                    paragraph(label+'：'+render_visible_text(f.get(key,'待补'),snapshot,evidence.get('evidence',[]),f.get('metric_refs',[]))).paragraph_format.keep_with_next=True
                paragraph('预期证据：'+'、'.join(render_visible_text(x,snapshot,evidence.get('evidence',[]),f.get('metric_refs',[])) for x in f.get('expected_evidence',[])))
            paragraph('任务须确认后才可模拟送达；送达不等于整改完成。人工归因0—5分、可读性及版式均待真人评审。')
            for c in snapshot['capabilities']:
                if c['status']!='available':paragraph(c['name']+'：'+str(c.get('reason')))
    paragraph('证据来源')
    for e in evidence.get('evidence',[]):paragraph((e.get('title') or '合成知识条目')+' · '+e.get('location','独立模拟条目'))
    style_reader(doc);output=Path(output);output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)
    check=verify(output,snapshot,benchmark)
    if check['status']!='PASS':raise ValueError('REPORT_CONTRACT_FAILED')
    return {'status':'PASS','scope':'file_generation','path':str(output),'sha256':hashlib.sha256(output.read_bytes()).hexdigest(),'verification':check}
