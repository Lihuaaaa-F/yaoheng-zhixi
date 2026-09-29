"""Report contracts are tested with independently constructed documents."""
from decimal import Decimal
from zipfile import ZipFile
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm
from pharma import reports


def test_toc_tab_matches_printable_width_in_twips():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)
    title = doc.add_paragraph('目录')
    head = doc.add_paragraph('一、封面与基本信息')
    reports._insert_native_toc(doc, title, [(head, head.text)])
    tab = doc.element.body.find('.//'+qn('w:tab'))
    expected = round((section.page_width-section.left_margin-section.right_margin)/635)
    assert abs(int(tab.get(qn('w:pos')))-expected) <= 1


def test_installed_template_verifies_its_own_numeric_binding_contract(tmp_path, monkeypatch):
    doc = Document()
    for heading in ['一、封面与基本信息','二、总成本概览','三、成本要素明细分析','四、重点产品专项分析','五、对标分析','六、总结与建议']:
        doc.add_paragraph(heading)
    placeholders=[]
    for i in range(20):
        p=doc.add_paragraph('金额100.00')
        mark=OxmlElement('w:bookmarkStart');mark.set(qn('w:name'),f'm{i}');mark.set(qn('w:id'),str(i));p._p.insert(0,mark)
        placeholders.append({'field':'total','original':'金额','context':'金额{{金额}}','marker':f'm{i}'})
    doc.add_paragraph('1.00')
    for _ in range(11):doc.add_table(rows=1,cols=1)
    path=tmp_path/'fixture.docx';doc.save(path)
    with ZipFile(path,'a') as z:z.writestr('word/media/fixture.png',b'fixture')
    monkeypatch.setattr(reports,'build_bindings',lambda *args:{'total':'100.00'})
    snapshot={'analysis_type':'monthly','metrics':{'unit_cost':{'value':'1'},'total_cost':{'value':'100'}}}
    result=reports.verify_docx(path,snapshot,placeholders=placeholders)
    assert result['status']=='PASS',result
    assert result['numeric_bindings_checked']==20
    assert result['expected_bindings']==20
    assessed=reports.assess_report({'docx':{'verification':result}})
    assert assessed['calculation_consistency']['status']=='PASS'
    placeholders[0]['context']='错误{{金额}}'
    assert reports.verify_docx(path,snapshot,placeholders=placeholders)['status']=='FAIL'


def test_key_conclusions_preserve_decimal_rounding_and_missing_deltas():
    snapshot={'metrics':{'unit_cost':{'value':'1'},'total_cost':{'value':'100'},'quantity':{'value':'100'}},
              'period_changes':{'unit_cost':{'mom':{'rate':'-0.405'}}},
              'elements':[{'name':'缺失要素','unit_delta':None}]}
    lines=reports._key_conclusion_lines(snapshot,{})
    assert '下降0.41%' in lines[0]
    assert not any('缺失要素 +0.00' in line for line in lines)


def test_incomplete_waterfall_is_not_drawn_as_zero(tmp_path):
    doc=Document();doc.add_paragraph('2.2 成本结构')
    snapshot={'product':'合成品','factory':'测试厂','month':'2030-06','analysis_type':'monthly',
              'trend':[], 'metrics':{'unit_cost':{'value':'10'}},
              'comparison':{'mom':{'base':'9'}}, 'elements':[{'name':'缺失要素','unit':None,'unit_delta':None}]}
    reports.add_reader_charts(doc,snapshot,None,{},tmp_path/'report.docx')
    assert not (tmp_path/'waterfall.png').exists()
    doc=Document()
    with pytest.raises(ValueError,match='REPORT_TEMPLATE_ANCHOR_MISSING:2.2'):
        reports.add_reader_charts(doc,snapshot,None,{},tmp_path/'report.docx')
