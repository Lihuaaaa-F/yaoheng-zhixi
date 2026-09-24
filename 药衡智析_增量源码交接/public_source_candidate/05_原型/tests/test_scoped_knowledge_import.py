"""Knowledge uploads preserve original locations and cannot cross enterprise scopes."""
import hashlib
import json

import pytest

from pharma import config, data_import, import_pipeline, knowledge


@pytest.fixture
def runtime(tmp_path, monkeypatch):
    monkeypatch.setattr(config, 'RUNTIME', tmp_path)
    monkeypatch.setattr(knowledge, 'RUNTIME', tmp_path)
    monkeypatch.setattr(data_import, 'RUNTIME', tmp_path)
    monkeypatch.setattr(data_import, 'IMPORTS_ROOT', tmp_path / 'imports')
    monkeypatch.setattr(data_import, 'IMPORT_DB', tmp_path / 'imports' / 'imports.sqlite3')
    monkeypatch.setattr(data_import, '_imports_db_ready', False)
    return tmp_path


def _publish(filename, payload, context='generic_manufacturing:enterprise-a'):
    record = data_import.create_upload('knowledge', filename, payload, 'enterprise')
    parsed = data_import.publish_knowledge(record, {'context_id': context})
    path = data_import.write_knowledge_source(record)
    return record, parsed, path


def test_pdf_page_and_source_link_survive_publication_and_retrieval(runtime):
    import fitz
    document = fitz.open()
    document.new_page().insert_text((36, 36), 'Equipment maintenance changed the production cycle. Inspect the maintenance log.')
    document.new_page().insert_text((36, 36), 'Packaging material purchase changed. Inspect the purchase invoices and yield records.')
    payload = document.tobytes()
    document.close()
    record, parsed, path = _publish('maintenance.pdf', payload)
    assert parsed['pages'] == 2
    blocks = list(knowledge.parse_document(path))
    assert [b['page'] for b in blocks] == [1, 2]
    kb = knowledge.Knowledge(root=runtime / 'isolated-index', context={'industry_id': 'generic_manufacturing', 'enterprise_id': 'enterprise-a'},
                             source_files=[path], include_uploads=False, vector_enabled=False)
    evidence = kb.search('Packaging purchase', mode='bm25')['evidence']
    assert any(item['page'] == 2 and item['source'] == 'maintenance.pdf' for item in evidence)
    assert all(item['source_id'] == parsed['source_id'] for item in evidence)
    source = data_import.resolve_knowledge_source(parsed['source_id'], parsed['context_id'])
    assert source['path'].read_bytes() == payload
    assert source['filename'] == record['filename']


def test_same_filename_revisions_keep_old_links_and_change_only_own_scope(runtime):
    context = 'generic_manufacturing:enterprise-a'
    other = 'generic_manufacturing:enterprise-b'
    before = knowledge.scoped_knowledge_snapshot(context, 'base')
    first, first_result, first_path = _publish('设备.txt', ('旧设备规程要求核对设备工时和采购记录。' * 5).encode(), context)
    version_a = knowledge.scoped_knowledge_snapshot(context, 'base')
    second, second_result, second_path = _publish('设备.txt', ('新设备规程要求检查加工损耗和质量记录。' * 5).encode(), context)
    assert before != version_a != knowledge.scoped_knowledge_snapshot(context, 'base')
    assert knowledge.scoped_knowledge_snapshot(other, 'base') == 'base'
    assert first_path != second_path and first_path.exists() and second_path.exists()
    registry = knowledge.knowledge_registry(context)
    assert registry['active']['enterprise:设备.txt'] == second_result['source_id']
    assert set(registry['versions']) == {first_result['source_id'], second_result['source_id']}
    assert knowledge.uploaded_knowledge_sources(other) == []
    old_source = data_import.resolve_knowledge_source(first_result['source_id'], context)
    assert old_source['path'].read_bytes() == data_import.original_path(first).read_bytes()
    assert json.loads(second_path.read_text())[0]['source'] == '设备.txt'
    with pytest.raises(KeyError):
        data_import.resolve_knowledge_source('../../etc/passwd', context)
    data_import.mark_import_status(data_import.get_import(first['id']), 'PARSE_FAILED')
    with pytest.raises(ValueError, match='IMPORT_ALREADY_REFERENCED'):
        data_import.delete_import(first['id'])
    assert old_source['path'].exists()


def test_txt_encoding_and_docx_paragraph_table_locations_are_real(runtime):
    import io
    from docx import Document
    _, _, text_path = _publish('说明.txt', ('设备检修后需要核查实际产量与成本归集口径。' * 4).encode('gb18030'))
    assert '设备检修' in json.loads(text_path.read_text())[0]['text']
    doc = Document()
    doc.add_paragraph('计量盘维护后需要结合维修工单确认设备状态，不能仅凭成本变化认定因果。')
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = '维修工单与产量记录需要共同确认业务原因'
    doc.add_paragraph('复核报告应保存采购记录、设备运行时长与质量检查结果。')
    doc.add_paragraph('负责人：张某')
    blob = io.BytesIO()
    doc.save(blob)
    _, result, path = _publish('设备.docx', blob.getvalue())
    blocks = json.loads(path.read_text())
    assert [b['location'] for b in blocks] == ['段落1', '表格1', '段落2', '段落3']
    assert all(b['page'] is None for b in blocks) and result['pages'] == 0
    kb = knowledge.Knowledge(root=runtime / 'docx-index', source_files=[path], vector_enabled=False)
    assert not kb.build()['failures']


def test_failed_index_build_never_activates_new_version(runtime, monkeypatch):
    from pharma import industry
    from pharma.jobs import JobStore
    context_id = 'generic_manufacturing:enterprise-a'
    _publish('运行规程.txt', ('已有设备运行规程要求核查业务记录。' * 5).encode(), context_id)
    registry_before = knowledge.knowledge_registry(context_id)
    base = runtime / 'knowledge.json'
    base.write_text('[]')
    context = industry.AnalysisContext(enterprise_id='enterprise-a', industry_id='generic_manufacturing',
        dataset_id='test', industry_version='1', policy_version='1', data_snapshot='data',
        knowledge_snapshot=knowledge.scoped_knowledge_snapshot(context_id, hashlib.sha256(base.read_bytes()).hexdigest()),
        template_version='1', formula_version='1')
    monkeypatch.setattr(industry, 'resolve_context', lambda _: context)
    monkeypatch.setattr(industry, 'knowledge_entry_for_context', lambda _: base)

    class BrokenIndex:
        def __init__(self, **kwargs):
            pass

        def build(self, **kwargs):
            return {'status': 'FAILED', 'failures': [{'reason': 'controlled index failure'}]}

    monkeypatch.setattr(knowledge, 'Knowledge', BrokenIndex)
    record = data_import.create_upload('knowledge', '运行规程.txt', ('更新规程仍需人工检查全部维修与质量记录。' * 5).encode(), 'enterprise')
    store = JobStore(runtime / 'jobs.sqlite3')
    job = store.enqueue('kb', {'import_ids': [record['id']], 'context_id': context_id})
    import_pipeline.run_kb_build(store, store.get(job['id']))
    assert store.get(job['id'])['status'] == 'FAILED'
    assert knowledge.knowledge_registry(context_id) == registry_before
    assert data_import.original_path(record).exists()


def test_existing_scoped_index_rebuilds_after_source_or_embedding_change(runtime):
    source = runtime / 'scoped-source.txt'
    source.write_text('维护规程要求核实设备工时与实际成本归集，证据不足时不能认定因果。')
    model_dir = runtime / 'model-assets'
    model_dir.mkdir()
    (model_dir / 'tokenizer.json').write_text('{"version": "one"}')
    kb = knowledge.Knowledge(root=runtime / 'index', source_files=[source], vector_enabled=False)
    kb.model_dir = model_dir
    first = kb.search('维护规程', mode='bm25')
    source.write_text('更新规程要求核实采购合同与设备维护记录，证据不足时不能认定因果。')
    second = kb.search('采购合同', mode='bm25')
    assert first['knowledge_version'] != second['knowledge_version']
    assert '采购合同' in second['evidence'][0]['text']
    (model_dir / 'tokenizer.json').write_text('{"version": "two changed"}')
    third = kb.search('采购合同', mode='bm25')
    assert second['knowledge_version'] != third['knowledge_version']


def test_mixed_scanned_pdf_is_not_silently_published_as_complete(runtime):
    import io
    import fitz
    from PIL import Image
    bitmap = io.BytesIO()
    Image.new('RGB', (20, 20), 'gray').save(bitmap, format='PNG')
    document = fitz.open()
    document.new_page().insert_text((36, 36), 'The text page has equipment maintenance records and operating instructions.')
    document.new_page().insert_image(fitz.Rect(0, 0, 100, 100), stream=bitmap.getvalue())
    payload = document.tobytes()
    document.close()
    record = data_import.create_upload('knowledge', 'mixed.pdf', payload, 'enterprise')
    with pytest.raises(ValueError, match='KNOWLEDGE_OCR_REQUIRED.*2'):
        data_import.publish_knowledge(record, {'context_id': 'generic_manufacturing:enterprise-a'})
    assert data_import.original_path(record).read_bytes() == payload
    assert knowledge.knowledge_registry('generic_manufacturing:enterprise-a')['active'] == {}


def test_workspace_inherits_verified_knowledge_without_retagging_or_cross_enterprise_reads(runtime, monkeypatch):
    current = 'generic_manufacturing:imp-workspace'
    former = ['generic_manufacturing:batch-a', 'generic_manufacturing:batch-b']
    first, first_meta, first_path = _publish('设备规程.txt', ('设备维护核查需要同时确认维修工单、生产工时与成本归集记录。' * 3).encode(), former[0])
    second, second_meta, second_path = _publish('采购规程.txt', ('采购价格核查需要同时确认采购订单、验收入库与付款凭证。' * 3).encode(), former[1])
    _, private_meta, private_path = _publish('其他企业.txt', ('其他企业机密规程不得进入当前企业的知识检索和报告内容。' * 3).encode(), 'generic_manufacturing:unrelated')
    _publish('比赛资料.txt', ('比赛资料与新企业不是同一个主体，不得自动混入新企业证据。' * 3).encode(), 'pharmaceutical:competition')
    originals = {p: p.read_bytes() for p in (first_path, second_path, private_path)}
    before = knowledge.scoped_knowledge_snapshot(current, 'base')
    monkeypatch.setattr(data_import, '_workspace_manifest', lambda: {
        'context_id': current, 'industry_id': 'generic_manufacturing', 'legacy_context_ids': former})
    assert set(knowledge.uploaded_knowledge_sources(current)) == {first_path, second_path}
    assert knowledge.scoped_knowledge_snapshot(current, 'base') != before
    assert knowledge.uploaded_knowledge_sources(former[0]) == [first_path]
    assert knowledge.registered_knowledge_source(second_meta['source_id'], former[0]) is None
    assert knowledge.registered_knowledge_source(private_meta['source_id'], current) is None
    assert knowledge.registered_knowledge_source(first_meta['source_id'], current)['source_context_id'] == former[0]
    resolved = data_import.resolve_knowledge_source(first_meta['source_id'], current)
    assert resolved['filename'] == first['filename']
    assert resolved['path'].read_bytes() == data_import.original_path(first).read_bytes()

    empty_base = runtime / 'workspace_knowledge.json'
    empty_base.write_text('[]')
    kb = knowledge.Knowledge(context={'industry_id': 'generic_manufacturing', 'enterprise_id': 'imp-workspace'},
                             source_files=[empty_base], vector_enabled=False)
    result = kb.search('设备维护 采购价格', mode='bm25')
    assert {row['source_id'] for row in result['evidence']} == {first_meta['source_id'], second_meta['source_id']}
    assert {row['source_context_id'] for row in result['evidence']} == set(former)
    assert all(row['analysis_context']['enterprise_id'] == 'imp-workspace' for row in result['evidence'])
    assert {p: p.read_bytes() for p in originals} == originals


def test_workspace_new_document_version_supersedes_inherited_active_but_retains_old_link(runtime, monkeypatch):
    current = 'generic_manufacturing:imp-workspace'
    former = 'generic_manufacturing:batch-a'
    _, old_meta, old_path = _publish('设备规程.txt', ('旧设备规程需要核对设备台账与历次检修工单。' * 3).encode(), former)
    monkeypatch.setattr(data_import, '_workspace_manifest', lambda: {
        'context_id': current, 'industry_id': 'generic_manufacturing', 'legacy_context_ids': [former]})
    before = knowledge.scoped_knowledge_snapshot(current, 'base')
    _, new_meta, new_path = _publish('设备规程.txt', ('更新的设备规程需要核对设备能耗、检修工单及产量记录。' * 3).encode(), current)
    assert knowledge.scoped_knowledge_snapshot(current, 'base') != before
    assert knowledge.uploaded_knowledge_sources(current) == [new_path]
    assert knowledge.uploaded_knowledge_sources(former) == [old_path]
    assert knowledge.registered_knowledge_source(old_meta['source_id'], current)['source_context_id'] == former
    assert knowledge.registered_knowledge_source(new_meta['source_id'], current)['source_context_id'] == current
    assert '旧设备规程' in data_import.resolve_knowledge_source(old_meta['source_id'], current)['path'].read_text()
    assert old_path.exists()


@pytest.mark.parametrize('foreign_scope', ['pharmaceutical:competition', 'mechanical_demo:enterprise-a'])
def test_workspace_manifest_cannot_extend_knowledge_to_contest_or_other_industry(runtime, monkeypatch, foreign_scope):
    current = 'generic_manufacturing:imp-workspace'
    monkeypatch.setattr(data_import, '_workspace_manifest', lambda: {
        'context_id': current, 'industry_id': 'generic_manufacturing', 'legacy_context_ids': [foreign_scope]})
    with pytest.raises(ValueError, match='INVALID_WORKSPACE_KNOWLEDGE_BINDING'):
        knowledge.uploaded_knowledge_sources(current)
    assert knowledge.knowledge_scope_ids('generic_manufacturing:unrelated') == ['generic_manufacturing:unrelated']
