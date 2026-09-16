from xml.etree import ElementTree as ET
from pharma.reports import replace_text_nodes, normalize_template

def test_cross_run_format_and_ambiguous_template_bindings(tmp_path):
    p = ET.fromstring('<p><t>{{人</t><t>工环比}}</t><t>%</t></p>')
    replace_text_nodes(list(p.iter('t')), {'人工环比': '2.15'})
    assert ''.join(t.text or '' for t in p.iter('t')) == '2.15%'
    result = normalize_template(tmp_path / 'working.docx', tmp_path / 'map.json')
    entries = result['placeholders']
    assert result['original_unique'] == 101
    assert result['original_occurrences'] == 108
    assert {r['field'] for r in entries if r['original'] == '人工环比'} == {'人工环比_金额','人工环比_比例'}
    assert {r['field'] for r in entries if r['original'] == '制造费用环比'} == {'制造费用环比_金额','制造费用环比_比例'}

def test_pdf_failure_preserves_docx(tmp_path):
    from pharma.reports import convert_pdf
    p=tmp_path/'keep.docx';p.write_bytes(b'fixture source')
    r=convert_pdf(p,converter='/no-such-converter')
    assert r['status']=='FAILED'
    assert p.read_bytes()==b'fixture source'

def test_document_gate_rejects_corrupted_numeric_cell(tmp_path):
    from pharma.reports import render_docx,verify_docx
    from pharma.metrics import analyze,benchmark
    from docx import Document
    s=analyze('中药一厂','银黄口服液','2026-05')
    p=tmp_path/'good.docx'
    render_docx(s,{'status':'DEGRADED','findings':[]},{'evidence':[]},p,benchmark(s['product'],s['month']))
    d=Document(p)
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text=='11.21':
                    for run in cell.paragraphs[0].runs:
                        if run.text=='11.21':run.text='999999.00'
    bad=tmp_path/'corrupt.docx';d.save(bad)
    check=verify_docx(bad,s)
    assert check['status']=='FAIL'
    assert check['numeric_binding_failures']

def test_missing_previous_period_total_contribution_is_na(tmp_path):
    from pharma.reports import build_bindings
    from pharma.metrics import analyze
    values=build_bindings(analyze('中药一厂','银黄口服液','2026-01'),{})
    assert values['合计贡献度'].startswith('N/A')


def test_indirect_labor_official_alias_is_bound():
    from pharma.reports import build_bindings
    from pharma.metrics import analyze
    s=analyze('中药一厂','六味地黄胶囊','2026-03')
    v=build_bindings(s,{})
    assert v['本月间接人工']=='0.57'
    assert v['上月间接人工']=='0.59'
